#!/usr/bin/env python3
"""Generate a premium client-ready Restaurant OS PowerPoint deck."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

# ── Design tokens (aligned to RR NOVA / pine–gold brand) ───────────────────────
C = {
    "bg_dark": RGBColor(0x0F, 0x1C, 0x18),
    "bg_dark2": RGBColor(0x1A, 0x2F, 0x28),
    "bg_panel": RGBColor(0x2A, 0x3F, 0x37),
    "bg_panel2": RGBColor(0x32, 0x4D, 0x42),
    "pine": RGBColor(0x4D, 0x77, 0x64),
    "pine_deep": RGBColor(0x3C, 0x5F, 0x50),
    "gold": RGBColor(0xC9, 0xA9, 0x62),
    "gold_soft": RGBColor(0xE2, 0xD4, 0xA8),
    "cream": RGBColor(0xF8, 0xF5, 0xF0),
    "cream2": RGBColor(0xFD, 0xFC, 0xFA),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "ink": RGBColor(0x1A, 0x2F, 0x28),
    "muted": RGBColor(0x6A, 0x94, 0x80),
    "muted_dark": RGBColor(0x97, 0xB8, 0xA9),
    "line": RGBColor(0xE4, 0xDC, 0xD0),
    "line_dark": RGBColor(0x3C, 0x5F, 0x50),
    "success": RGBColor(0x3D, 0xA3, 0x72),
    "warn": RGBColor(0xD4, 0xA0, 0x4A),
    "danger": RGBColor(0xC4, 0x5C, 0x4A),
    "blue": RGBColor(0x4A, 0x7F, 0xA8),
    "card_light": RGBColor(0xFF, 0xFF, 0xFF),
    "ui_bg": RGBColor(0xF2, 0xF6, 0xF4),
    "ui_sidebar": RGBColor(0x1A, 0x2F, 0x28),
}

FONT_SANS = "Avenir Next"
FONT_DISPLAY = "Georgia"
FONT_FALLBACK = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.55)


def _set_run_font(run, size, bold=False, color=None, font=FONT_SANS, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color


def _fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _line(shape, color: RGBColor | None = None, width_pt: float = 0.75):
    if color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color
        shape.line.width = Pt(width_pt)


def _set_shape_shadow(shape, soft=True):
    """Add a subtle drop shadow via OOXML."""
    spPr = shape._element.spPr
    effectLst = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    if effectLst is None:
        effectLst = OxmlElement("a:effectLst")
        spPr.append(effectLst)
    for child in list(effectLst):
        effectLst.remove(child)
    outer = OxmlElement("a:outerShdw")
    outer.set("blurRad", "50800" if soft else "76200")
    outer.set("dist", "38100" if soft else "50800")
    outer.set("dir", "2700000")
    outer.set("algn", "tl")
    outer.set("rotWithShape", "0")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", "0F1C18")
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "18000" if soft else "28000")
    srgb.append(alpha)
    outer.append(srgb)
    effectLst.append(outer)


def add_rect(slide, left, top, width, height, fill, line=None, radius=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    _fill(shp, fill)
    _line(shp, line)
    if radius is not None and hasattr(shp, "adjustments") and len(shp.adjustments) > 0:
        # 0–1 adj; smaller = tighter radius relative to size
        shp.adjustments[0] = radius
    if shadow:
        _set_shape_shadow(shp)
    return shp


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=14,
    bold=False,
    color=C["ink"],
    font=FONT_SANS,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    italic=False,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, bold=bold, color=color, font=font, italic=italic)
    return box


def add_paras(
    slide,
    left,
    top,
    width,
    height,
    lines: list[tuple[str, dict]],
):
    """lines: list of (text, style kwargs)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, style) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = style.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(style.get("space_before", 0))
        p.space_after = Pt(style.get("space_after", 4))
        run = p.add_run()
        run.text = text
        _set_run_font(
            run,
            style.get("size", 14),
            bold=style.get("bold", False),
            color=style.get("color", C["ink"]),
            font=style.get("font", FONT_SANS),
            italic=style.get("italic", False),
        )
    return box


def solid_bg(slide, color):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, color)


def accent_bar(slide, dark=True):
    add_rect(slide, 0, 0, Inches(0.08), SLIDE_H, C["gold"])


def brand_chip(slide, left, top, dark=True):
    add_rect(slide, left, top, Inches(0.34), Inches(0.34), C["gold"], radius=0.2)
    add_text(
        slide,
        left,
        top + Inches(0.02),
        Inches(0.34),
        Inches(0.3),
        "R",
        size=12,
        bold=True,
        color=C["bg_dark"],
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        left + Inches(0.42),
        top + Inches(0.02),
        Inches(3),
        Inches(0.32),
        "Restaurant OS",
        size=13,
        bold=True,
        color=C["white"] if dark else C["ink"],
        font=FONT_DISPLAY,
    )


def eyebrow(slide, left, top, text, dark=True):
    add_text(
        slide,
        left,
        top,
        Inches(6),
        Inches(0.28),
        text.upper(),
        size=10,
        bold=True,
        color=C["gold"] if dark else C["pine"],
    )


def section_title(slide, left, top, text, dark=True, width=Inches(7.5), size=32):
    add_text(
        slide,
        left,
        top,
        width,
        Inches(0.7),
        text,
        size=size,
        bold=True,
        color=C["white"] if dark else C["ink"],
        font=FONT_DISPLAY,
    )


def body(slide, left, top, width, text, dark=True, size=15):
    add_text(
        slide,
        left,
        top,
        width,
        Inches(1.2),
        text,
        size=size,
        color=C["muted_dark"] if dark else C["pine_deep"],
    )


def bullet_block(slide, left, top, width, items, dark=False, size=14):
    box = slide.shapes.add_textbox(left, top, width, Inches(3.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_before = Pt(6)
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = f"•  {item}"
        _set_run_font(
            run,
            size,
            color=C["muted_dark"] if dark else C["ink"],
        )
    return box


def meta_pills(slide, left, top, pills: list[tuple[str, str]], dark=False):
    """pills: (label, value) e.g. Pain / Benefit."""
    x = left
    for label, value in pills:
        w = Inches(2.55)
        bg = C["bg_panel"] if dark else C["ui_bg"]
        add_rect(slide, x, top, w, Inches(0.72), bg, radius=0.15)
        add_text(slide, x + Inches(0.14), top + Inches(0.08), w - Inches(0.2), Inches(0.22), label.upper(), size=8, bold=True, color=C["gold"] if dark else C["pine"])
        add_text(slide, x + Inches(0.14), top + Inches(0.32), w - Inches(0.2), Inches(0.32), value, size=11, bold=True, color=C["white"] if dark else C["ink"])
        x += w + Inches(0.12)


def card(slide, left, top, width, height, fill=None, dark=False):
    fill = fill or (C["bg_panel"] if dark else C["white"])
    shp = add_rect(slide, left, top, width, height, fill, line=None if dark else C["line"], radius=0.08, shadow=not dark)
    return shp


# ── UI mockup helpers ───────────────────────────────────────────────────────

def browser_chrome(slide, left, top, width, height, url="app.restaurantos.nz", dark_frame=True):
    """Draw a product window frame and return content origin (x, y, w, h)."""
    frame = add_rect(slide, left, top, width, height, C["bg_dark2"] if dark_frame else C["white"], radius=0.06, shadow=True)
    # title bar
    bar_h = Inches(0.36)
    add_rect(slide, left, top, width, bar_h, C["bg_dark"] if dark_frame else C["ui_bg"], radius=0.0)
    # traffic lights
    for i, col in enumerate([C["danger"], C["warn"], C["success"]]):
        add_rect(slide, left + Inches(0.14) + Inches(i * 0.18), top + Inches(0.12), Inches(0.1), Inches(0.1), col, radius=0.5)
    # url pill
    add_rect(slide, left + Inches(1.1), top + Inches(0.08), width - Inches(1.4), Inches(0.2), C["bg_panel"] if dark_frame else C["cream"], radius=0.3)
    add_text(slide, left + Inches(1.2), top + Inches(0.07), width - Inches(1.6), Inches(0.2), url, size=8, color=C["muted_dark"] if dark_frame else C["muted"], align=PP_ALIGN.LEFT)
    content_top = top + bar_h
    content_h = height - bar_h
    return left, content_top, width, content_h


def dash_shell(slide, left, top, width, height, active="Overview", rows=None):
    """Mini dashboard shell with sidebar + content area. Returns content rect."""
    l, t, w, h = browser_chrome(slide, left, top, width, height, url="app.restaurantos.nz/dashboard")
    side_w = Inches(1.35)
    add_rect(slide, l, t, side_w, h, C["ui_sidebar"])
    add_text(slide, l + Inches(0.1), t + Inches(0.12), side_w - Inches(0.15), Inches(0.25), "Harbour Kitchen", size=8, bold=True, color=C["gold"])
    nav = ["Overview", "Orders", "Menu", "Website", "QR", "Bookings", "Guests", "Analytics", "AI", "Team"]
    for i, name in enumerate(nav):
        y = t + Inches(0.45) + Inches(i * 0.28)
        if name == active:
            add_rect(slide, l + Inches(0.08), y - Inches(0.04), side_w - Inches(0.16), Inches(0.26), C["pine_deep"], radius=0.15)
            col = C["gold"]
        else:
            col = C["muted_dark"]
        add_text(slide, l + Inches(0.16), y, side_w - Inches(0.25), Inches(0.22), name, size=8, color=col)
    content_l = l + side_w
    content_w = w - side_w
    add_rect(slide, content_l, t, content_w, h, C["ui_bg"])
    return content_l, t, content_w, h


def mock_kpi_row(slide, left, top, width, metrics):
    gap = Inches(0.1)
    n = len(metrics)
    card_w = (width - gap * (n - 1)) / n
    for i, (label, value, delta) in enumerate(metrics):
        x = left + (card_w + gap) * i
        add_rect(slide, x, top, card_w, Inches(0.78), C["white"], radius=0.1)
        add_text(slide, x + Inches(0.1), top + Inches(0.08), card_w - Inches(0.15), Inches(0.18), label, size=7, color=C["muted"])
        add_text(slide, x + Inches(0.1), top + Inches(0.28), card_w - Inches(0.15), Inches(0.28), value, size=14, bold=True, color=C["ink"])
        add_text(slide, x + Inches(0.1), top + Inches(0.55), card_w - Inches(0.15), Inches(0.18), delta, size=7, color=C["success"])


def mock_table(slide, left, top, width, headers, rows, row_h=Inches(0.28)):
    add_rect(slide, left, top, width, Inches(0.28) + row_h * len(rows), C["white"], radius=0.08)
    col_w = width / len(headers)
    for i, h in enumerate(headers):
        add_text(slide, left + col_w * i + Inches(0.08), top + Inches(0.04), col_w - Inches(0.1), Inches(0.2), h, size=7, bold=True, color=C["muted"])
    for r, row in enumerate(rows):
        y = top + Inches(0.28) + row_h * r
        if r % 2 == 0:
            add_rect(slide, left, y, width, row_h, C["cream2"])
        for i, cell in enumerate(row):
            add_text(slide, left + col_w * i + Inches(0.08), y + Inches(0.04), col_w - Inches(0.1), row_h - Inches(0.04), cell, size=8, color=C["ink"])


def feature_slide(
    prs,
    *,
    title: str,
    eyebrow_text: str,
    bullets: list[str],
    pain: str,
    benefit: str,
    mock_fn,
    dark=False,
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(slide, C["bg_dark"] if dark else C["cream"])
    accent_bar(slide)
    brand_chip(slide, MARGIN, Inches(0.28), dark=dark)
    eyebrow(slide, MARGIN, Inches(0.85), eyebrow_text, dark=dark)
    section_title(slide, MARGIN, Inches(1.1), title, dark=dark, width=Inches(6.2), size=28)
    bullet_block(slide, MARGIN, Inches(1.9), Inches(5.6), bullets, dark=dark, size=13)
    meta_pills(
        slide,
        MARGIN,
        Inches(5.55),
        [("Pain solved", pain), ("Business benefit", benefit)],
        dark=dark,
    )
    # mockup on right
    mock_fn(slide, Inches(6.55), Inches(0.85), Inches(6.2), Inches(5.9))
    return slide


# ── Specific mockups ────────────────────────────────────────────────────────

def mock_dashboard(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Overview")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.12), Inches(3), Inches(0.25), "Today · Auckland NZ", size=10, bold=True, color=C["ink"])
    mock_kpi_row(
        slide,
        cl + Inches(0.15),
        ct + Inches(0.45),
        cw - Inches(0.3),
        [
            ("Revenue", "NZ$4,820", "+18% vs avg"),
            ("Orders", "86", "+12 open"),
            ("Covers", "142", "92% booked"),
            ("QR scans", "218", "Table 5 hot"),
        ],
    )
    add_rect(slide, cl + Inches(0.15), ct + Inches(1.4), cw - Inches(0.3), Inches(1.7), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.28), ct + Inches(1.5), Inches(3), Inches(0.22), "Live orders", size=9, bold=True, color=C["ink"])
    mock_table(
        slide,
        cl + Inches(0.2),
        ct + Inches(1.8),
        cw - Inches(0.4),
        ["Order", "Type", "Status", "Total"],
        [
            ["#1842", "Dine-in · T5", "Preparing", "$86.50"],
            ["#1841", "Pickup", "Ready", "$42.00"],
            ["#1840", "Delivery", "Out", "$67.20"],
            ["#1839", "Dine-in · T12", "Served", "$124.00"],
        ],
    )
    add_rect(slide, cl + Inches(0.15), ct + Inches(3.7), (cw - Inches(0.4)) / 2, Inches(1.35), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.28), ct + Inches(3.85), Inches(2.5), Inches(0.2), "Tonight's reservations", size=9, bold=True, color=C["ink"])
    add_text(slide, cl + Inches(0.28), ct + Inches(4.2), Inches(2.5), Inches(0.7), "7:00  Chen · 4\n7:30  Patel · 2\n8:00  Wilson · 6", size=9, color=C["pine_deep"])
    add_rect(slide, cl + Inches(0.15) + (cw - Inches(0.4)) / 2 + Inches(0.1), ct + Inches(3.7), (cw - Inches(0.4)) / 2, Inches(1.35), C["bg_dark2"], radius=0.08)
    add_text(slide, cl + Inches(0.3) + (cw - Inches(0.4)) / 2 + Inches(0.1), ct + Inches(3.85), Inches(2.4), Inches(0.2), "AI insight", size=9, bold=True, color=C["gold"])
    add_text(slide, cl + Inches(0.3) + (cw - Inches(0.4)) / 2 + Inches(0.1), ct + Inches(4.15), Inches(2.4), Inches(0.8), "Lamb shoulder sells out by 8pm Fridays. Suggest a limited add-on.", size=9, color=C["muted_dark"])


def mock_settings(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Website")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.12), Inches(4), Inches(0.25), "Restaurant profile & settings", size=11, bold=True, color=C["ink"])
    # form cards
    add_rect(slide, cl + Inches(0.15), ct + Inches(0.5), cw - Inches(0.3), Inches(1.5), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.3), ct + Inches(0.6), Inches(3), Inches(0.2), "Brand identity", size=9, bold=True, color=C["ink"])
    add_rect(slide, cl + Inches(0.3), ct + Inches(0.95), Inches(0.7), Inches(0.7), C["pine"], radius=0.12)
    add_text(slide, cl + Inches(0.45), ct + Inches(1.15), Inches(0.5), Inches(0.3), "HK", size=12, bold=True, color=C["gold"], align=PP_ALIGN.CENTER)
    add_text(slide, cl + Inches(1.2), ct + Inches(0.95), Inches(3.5), Inches(0.7), "Harbour Kitchen\nWynyard Quarter, Auckland\nTimezone: Pacific/Auckland · GST on", size=9, color=C["pine_deep"])
    fields = [
        ("Phone", "+64 9 555 0142"),
        ("Email", "hello@harbourkitchen.nz"),
        ("Currency", "NZD"),
        ("Cuisine", "Modern NZ · Seafood"),
    ]
    for i, (lab, val) in enumerate(fields):
        x = cl + Inches(0.15) + (i % 2) * ((cw - Inches(0.4)) / 2 + Inches(0.1))
        y = ct + Inches(2.2) + (i // 2) * Inches(0.85)
        add_rect(slide, x, y, (cw - Inches(0.4)) / 2, Inches(0.72), C["white"], radius=0.08)
        add_text(slide, x + Inches(0.14), y + Inches(0.1), Inches(2), Inches(0.18), lab, size=8, color=C["muted"])
        add_text(slide, x + Inches(0.14), y + Inches(0.32), Inches(2.5), Inches(0.28), val, size=11, bold=True, color=C["ink"])
    add_rect(slide, cl + Inches(0.15), ct + Inches(4.0), cw - Inches(0.3), Inches(0.9), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.3), ct + Inches(4.15), Inches(4), Inches(0.2), "Opening hours", size=9, bold=True, color=C["ink"])
    add_text(slide, cl + Inches(0.3), ct + Inches(4.45), Inches(5), Inches(0.35), "Mon–Thu 11:30–22:00 · Fri–Sat 11:30–23:00 · Sun 11:00–21:00", size=9, color=C["pine_deep"])


def mock_website(slide, left, top, width, height):
    l, t, w, h = browser_chrome(slide, left, top, width, height, url="harbourkitchen.nz", dark_frame=False)
    # hero
    add_rect(slide, l, t, w, h * 0.55, C["bg_dark2"])
    add_text(slide, l + Inches(0.3), t + Inches(0.25), w - Inches(0.5), Inches(0.25), "HARBOUR KITCHEN", size=10, bold=True, color=C["gold"])
    add_text(slide, l + Inches(0.3), t + Inches(0.6), w - Inches(0.6), Inches(0.7), "Waterfront dining\nin Wynyard Quarter", size=18, bold=True, color=C["white"], font=FONT_DISPLAY)
    add_text(slide, l + Inches(0.3), t + Inches(1.45), w - Inches(0.6), Inches(0.4), "Seasonal menus, harbour views, and effortless online booking.", size=9, color=C["muted_dark"])
    add_rect(slide, l + Inches(0.3), t + Inches(2.0), Inches(1.3), Inches(0.32), C["gold"], radius=0.2)
    add_text(slide, l + Inches(0.3), t + Inches(2.02), Inches(1.3), Inches(0.28), "Reserve a table", size=8, bold=True, color=C["bg_dark"], align=PP_ALIGN.CENTER)
    add_rect(slide, l + Inches(1.75), t + Inches(2.0), Inches(1.1), Inches(0.32), C["bg_panel"], radius=0.2)
    add_text(slide, l + Inches(1.75), t + Inches(2.02), Inches(1.1), Inches(0.28), "View menu", size=8, bold=True, color=C["white"], align=PP_ALIGN.CENTER)
    # lower sections
    add_rect(slide, l, t + h * 0.55, w, h * 0.45, C["cream"])
    for i, (title, body) in enumerate([
        ("Dinner", "Market fish, lamb, produce"),
        ("Wine", "NZ & coastal pairings"),
        ("Events", "Private dining suite"),
    ]):
        x = l + Inches(0.2) + i * ((w - Inches(0.5)) / 3 + Inches(0.08))
        add_rect(slide, x, t + h * 0.55 + Inches(0.2), (w - Inches(0.5)) / 3, Inches(1.5), C["white"], radius=0.08)
        add_text(slide, x + Inches(0.12), t + h * 0.55 + Inches(0.35), Inches(1.5), Inches(0.25), title, size=11, bold=True, color=C["ink"], font=FONT_DISPLAY)
        add_text(slide, x + Inches(0.12), t + h * 0.55 + Inches(0.7), Inches(1.5), Inches(0.6), body, size=8, color=C["pine_deep"])


def mock_menu(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Menu")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(3), Inches(0.25), "Dinner menu · Live", size=11, bold=True, color=C["ink"])
    add_rect(slide, cl + cw - Inches(1.4), ct + Inches(0.1), Inches(1.15), Inches(0.28), C["pine"], radius=0.2)
    add_text(slide, cl + cw - Inches(1.4), ct + Inches(0.12), Inches(1.15), Inches(0.24), "Publish", size=8, bold=True, color=C["white"], align=PP_ALIGN.CENTER)
    cats = ["Starters", "Mains", "Sides", "Drinks"]
    for i, c in enumerate(cats):
        x = cl + Inches(0.15) + i * Inches(0.95)
        bg = C["pine"] if i == 1 else C["white"]
        fg = C["white"] if i == 1 else C["ink"]
        add_rect(slide, x, ct + Inches(0.5), Inches(0.88), Inches(0.28), bg, radius=0.2)
        add_text(slide, x, ct + Inches(0.52), Inches(0.88), Inches(0.24), c, size=8, bold=True, color=fg, align=PP_ALIGN.CENTER)
    items = [
        ("Market Fish of the Day", "NZ$42", "On"),
        ("Slow-Cooked Lamb Shoulder", "NZ$48", "On"),
        ("Harbour Burger", "NZ$28", "On"),
        ("Charred Broccolini", "NZ$14", "On"),
        ("Espresso Martini", "NZ$18", "86'd"),
    ]
    for i, (name, price, flag) in enumerate(items):
        y = ct + Inches(1.0) + i * Inches(0.7)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.62), C["white"], radius=0.08)
        add_rect(slide, cl + Inches(0.25), y + Inches(0.1), Inches(0.42), Inches(0.42), C["ui_bg"], radius=0.1)
        add_text(slide, cl + Inches(0.85), y + Inches(0.12), Inches(3.2), Inches(0.22), name, size=10, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(0.85), y + Inches(0.35), Inches(2), Inches(0.2), "Modifiers · allergens · photo", size=7, color=C["muted"])
        add_text(slide, cl + cw - Inches(1.55), y + Inches(0.18), Inches(0.7), Inches(0.25), price, size=10, bold=True, color=C["ink"])
        col = C["danger"] if flag == "86'd" else C["success"]
        add_text(slide, cl + cw - Inches(0.85), y + Inches(0.2), Inches(0.55), Inches(0.22), flag, size=8, bold=True, color=col)


def mock_qr(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="QR")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Table QR codes · Floor plan", size=11, bold=True, color=C["ink"])
    # grid of tables
    for i in range(12):
        row, col = divmod(i, 4)
        x = cl + Inches(0.2) + col * Inches(1.35)
        y = ct + Inches(0.55) + row * Inches(1.35)
        add_rect(slide, x, y, Inches(1.2), Inches(1.2), C["white"], radius=0.1)
        # fake QR
        add_rect(slide, x + Inches(0.25), y + Inches(0.15), Inches(0.7), Inches(0.7), C["ink"], radius=0.05)
        add_rect(slide, x + Inches(0.35), y + Inches(0.25), Inches(0.5), Inches(0.5), C["white"])
        add_rect(slide, x + Inches(0.42), y + Inches(0.32), Inches(0.36), Inches(0.36), C["ink"])
        add_text(slide, x, y + Inches(0.9), Inches(1.2), Inches(0.22), f"Table {i + 1}", size=8, bold=True, color=C["ink"], align=PP_ALIGN.CENTER)


def mock_ordering(slide, left, top, width, height):
    l, t, w, h = browser_chrome(slide, left, top, width, height, url="harbourkitchen.nz/order", dark_frame=False)
    add_rect(slide, l, t, w, h, C["cream"])
    add_text(slide, l + Inches(0.2), t + Inches(0.15), Inches(4), Inches(0.25), "Order online · Pickup or delivery", size=11, bold=True, color=C["ink"])
    # mode tabs
    for i, (lab, on) in enumerate([("Pickup", True), ("Delivery", False), ("Dine-in", False)]):
        x = l + Inches(0.2) + i * Inches(1.15)
        add_rect(slide, x, t + Inches(0.5), Inches(1.05), Inches(0.3), C["pine"] if on else C["white"], radius=0.2)
        add_text(slide, x, t + Inches(0.52), Inches(1.05), Inches(0.26), lab, size=8, bold=True, color=C["white"] if on else C["ink"], align=PP_ALIGN.CENTER)
    cart = [
        ("Harbour Burger", "1 × $28.00"),
        ("Truffle Fries", "1 × $12.00"),
        ("Sparkling Water", "2 × $6.00"),
    ]
    for i, (n, p) in enumerate(cart):
        y = t + Inches(1.05) + i * Inches(0.7)
        add_rect(slide, l + Inches(0.2), y, w - Inches(0.4), Inches(0.6), C["white"], radius=0.08)
        add_text(slide, l + Inches(0.35), y + Inches(0.12), Inches(3), Inches(0.22), n, size=10, bold=True, color=C["ink"])
        add_text(slide, l + Inches(0.35), y + Inches(0.34), Inches(2), Inches(0.18), p, size=8, color=C["muted"])
    add_rect(slide, l + Inches(0.2), t + Inches(3.3), w - Inches(0.4), Inches(1.6), C["white"], radius=0.08)
    add_text(slide, l + Inches(0.35), t + Inches(3.45), Inches(3), Inches(0.22), "Order summary", size=9, bold=True, color=C["ink"])
    add_text(slide, l + Inches(0.35), t + Inches(3.8), Inches(4), Inches(0.7), "Subtotal  NZ$52.00\nGST      NZ$7.80\nTotal     NZ$59.80", size=10, color=C["pine_deep"])
    add_rect(slide, l + Inches(0.35), t + Inches(4.5), Inches(2.2), Inches(0.32), C["gold"], radius=0.2)
    add_text(slide, l + Inches(0.35), t + Inches(4.52), Inches(2.2), Inches(0.28), "Pay securely", size=9, bold=True, color=C["bg_dark"], align=PP_ALIGN.CENTER)


def mock_fulfillment(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Orders")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Pickup & delivery board", size=11, bold=True, color=C["ink"])
    cols = [
        ("Pickup queue", [("Chen", "Ready 5:40"), ("Ng", "Prep 5:55"), ("Lee", "New 6:05")], C["success"]),
        ("Out for delivery", [("Patel · CBD", "Driver 2"), ("Singh · Ponsonby", "Driver 1")], C["blue"]),
        ("Completed", [("Wilson", "5:20"), ("Ava", "5:05")], C["muted"]),
    ]
    col_w = (cw - Inches(0.5)) / 3
    for i, (title, items, accent) in enumerate(cols):
        x = cl + Inches(0.15) + i * (col_w + Inches(0.1))
        add_rect(slide, x, ct + Inches(0.5), col_w, ch - Inches(0.7), C["white"], radius=0.08)
        add_rect(slide, x, ct + Inches(0.5), col_w, Inches(0.08), accent)
        add_text(slide, x + Inches(0.12), ct + Inches(0.7), col_w - Inches(0.2), Inches(0.25), title, size=9, bold=True, color=C["ink"])
        for j, (name, meta) in enumerate(items):
            y = ct + Inches(1.15) + j * Inches(0.95)
            add_rect(slide, x + Inches(0.1), y, col_w - Inches(0.2), Inches(0.8), C["ui_bg"], radius=0.1)
            add_text(slide, x + Inches(0.2), y + Inches(0.15), col_w - Inches(0.35), Inches(0.22), name, size=10, bold=True, color=C["ink"])
            add_text(slide, x + Inches(0.2), y + Inches(0.42), col_w - Inches(0.35), Inches(0.22), meta, size=8, color=C["muted"])


def mock_kds(slide, left, top, width, height):
    l, t, w, h = browser_chrome(slide, left, top, width, height, url="app.restaurantos.nz/kitchen", dark_frame=True)
    add_rect(slide, l, t, w, h, C["bg_dark"])
    add_text(slide, l + Inches(0.2), t + Inches(0.12), Inches(4), Inches(0.25), "Kitchen display · Live", size=11, bold=True, color=C["gold"])
    tickets = [
        ("#1842 · T5", "2× Lamb · 1× Fish", "8:12", C["danger"]),
        ("#1841 · PU", "Burger · Fries", "4:05", C["warn"]),
        ("#1840 · DL", "Salad · Coffee", "2:40", C["success"]),
        ("#1838 · T12", "Burger ×2 · Broccolini", "1:15", C["pine"]),
    ]
    for i, (oid, items, age, accent) in enumerate(tickets):
        row, col = divmod(i, 2)
        x = l + Inches(0.2) + col * ((w - Inches(0.5)) / 2 + Inches(0.1))
        y = t + Inches(0.55) + row * Inches(2.2)
        tw = (w - Inches(0.5)) / 2
        add_rect(slide, x, y, tw, Inches(2.0), C["bg_panel"], radius=0.08)
        add_rect(slide, x, y, Inches(0.1), Inches(2.0), accent)
        add_text(slide, x + Inches(0.25), y + Inches(0.2), tw - Inches(0.4), Inches(0.25), oid, size=12, bold=True, color=C["white"])
        add_text(slide, x + Inches(0.25), y + Inches(0.55), tw - Inches(0.4), Inches(0.7), items, size=11, color=C["muted_dark"])
        add_text(slide, x + Inches(0.25), y + Inches(1.4), tw - Inches(0.4), Inches(0.25), f"{age} elapsed", size=10, bold=True, color=accent)
        add_rect(slide, x + Inches(0.25), y + Inches(1.7), Inches(1.2), Inches(0.22), accent, radius=0.2)
        add_text(slide, x + Inches(0.25), y + Inches(1.7), Inches(1.2), Inches(0.22), "Bump", size=8, bold=True, color=C["white"], align=PP_ALIGN.CENTER)


def mock_reservations(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Bookings")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Reservations · Fri 12 Sep", size=11, bold=True, color=C["ink"])
    mock_kpi_row(
        slide,
        cl + Inches(0.15),
        ct + Inches(0.45),
        cw - Inches(0.3),
        [("Booked", "68 covers", "Dinner"), ("Available", "24 covers", "Until 9pm"), ("No-shows", "2%", "This week"), ("Waitlist", "7 parties", "Live")],
    )
    mock_table(
        slide,
        cl + Inches(0.15),
        ct + Inches(1.45),
        cw - Inches(0.3),
        ["Time", "Guest", "Party", "Status"],
        [
            ["17:30", "Ava Chen", "4", "Confirmed"],
            ["18:00", "James Patel", "2", "Confirmed"],
            ["18:30", "Mia Wilson", "6", "Seated"],
            ["19:00", "Noah Lee", "3", "Pending"],
            ["19:30", "Sofia Ng", "2", "Confirmed"],
            ["20:00", "Private dining", "10", "Deposit paid"],
        ],
    )


def mock_waitlist(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Bookings")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Waitlist · Live floor", size=11, bold=True, color=C["ink"])
    parties = [
        ("Chen party of 4", "Quoted 25 min", "SMS ready", C["warn"]),
        ("Patel party of 2", "Quoted 15 min", "Notified", C["success"]),
        ("Wilson party of 6", "Quoted 40 min", "Waiting", C["blue"]),
        ("Lee party of 3", "Quoted 20 min", "SMS ready", C["warn"]),
    ]
    for i, (name, quote, status, col) in enumerate(parties):
        y = ct + Inches(0.55) + i * Inches(1.0)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.88), C["white"], radius=0.08)
        add_rect(slide, cl + Inches(0.15), y, Inches(0.1), Inches(0.88), col)
        add_text(slide, cl + Inches(0.4), y + Inches(0.15), Inches(3), Inches(0.25), name, size=12, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(0.4), y + Inches(0.45), Inches(3), Inches(0.25), quote, size=9, color=C["muted"])
        add_rect(slide, cl + cw - Inches(1.7), y + Inches(0.28), Inches(1.25), Inches(0.32), col, radius=0.2)
        add_text(slide, cl + cw - Inches(1.7), y + Inches(0.3), Inches(1.25), Inches(0.28), status, size=8, bold=True, color=C["white"], align=PP_ALIGN.CENTER)


def mock_crm(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Guests")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Customers · CRM", size=11, bold=True, color=C["ink"])
    mock_table(
        slide,
        cl + Inches(0.15),
        ct + Inches(0.5),
        cw - Inches(0.3),
        ["Guest", "Visits", "Spend", "Last"],
        [
            ["Ava Chen", "14", "$1,842", "Yesterday"],
            ["James Patel", "9", "$980", "3 days"],
            ["Mia Wilson", "7", "$1,210", "Last week"],
            ["Noah Lee", "5", "$640", "Today"],
            ["Sofia Ng", "11", "$1,540", "Fri"],
        ],
    )
    add_rect(slide, cl + Inches(0.15), ct + Inches(3.3), cw - Inches(0.3), Inches(1.7), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.3), ct + Inches(3.45), Inches(4), Inches(0.22), "Guest profile · Ava Chen", size=10, bold=True, color=C["ink"])
    add_text(slide, cl + Inches(0.3), ct + Inches(3.8), Inches(5), Inches(1.0), "Prefers window table · Allergic to shellfish\nAvg ticket NZ$132 · Loyalty Gold · Birthday 18 Oct\nChannels: email + SMS · Marketing OK", size=9, color=C["pine_deep"])


def mock_loyalty(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Guests")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Loyalty programme", size=11, bold=True, color=C["ink"])
    tiers = [("Bronze", "0–499 pts", C["pine"]), ("Silver", "500–999", C["muted"]), ("Gold", "1,000+", C["gold"])]
    for i, (n, r, col) in enumerate(tiers):
        x = cl + Inches(0.15) + i * ((cw - Inches(0.5)) / 3 + Inches(0.1))
        add_rect(slide, x, ct + Inches(0.5), (cw - Inches(0.5)) / 3, Inches(1.3), C["white"], radius=0.08)
        add_rect(slide, x + Inches(0.15), ct + Inches(0.7), Inches(0.35), Inches(0.35), col, radius=0.5)
        add_text(slide, x + Inches(0.6), ct + Inches(0.75), Inches(1.5), Inches(0.25), n, size=12, bold=True, color=C["ink"])
        add_text(slide, x + Inches(0.15), ct + Inches(1.25), Inches(1.8), Inches(0.3), r, size=9, color=C["muted"])
    add_rect(slide, cl + Inches(0.15), ct + Inches(2.05), cw - Inches(0.3), Inches(2.9), C["bg_dark2"], radius=0.08)
    add_text(slide, cl + Inches(0.35), ct + Inches(2.25), Inches(4), Inches(0.25), "Earn 1 pt per NZ$1 · Redeem rewards", size=11, bold=True, color=C["gold"])
    add_text(slide, cl + Inches(0.35), ct + Inches(2.7), Inches(5), Inches(1.8), "• Free dessert at 250 pts\n• $25 credit at 500 pts\n• Chef's table priority for Gold\n• Birthday reward auto-sent\n• Points visible on guest receipt & profile", size=11, color=C["muted_dark"])


def mock_offers(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Overview")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Offers & promotions", size=11, bold=True, color=C["ink"])
    offers = [
        ("Weekday Lunch Set", "Active", "Mon–Thu · 11:30–14:30", "NZ$39"),
        ("Happy Hour Wine", "Active", "Fri · 16:00–18:00", "30% off"),
        ("Winter Warmers", "Scheduled", "Starts 1 Jun", "Bundle"),
        ("Loyalty Double Points", "Draft", "Email segment: Gold", "2× pts"),
    ]
    for i, (n, st, when, deal) in enumerate(offers):
        y = ct + Inches(0.5) + i * Inches(1.05)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.92), C["white"], radius=0.08)
        add_text(slide, cl + Inches(0.3), y + Inches(0.15), Inches(3.5), Inches(0.25), n, size=12, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(0.3), y + Inches(0.45), Inches(3.5), Inches(0.25), when, size=9, color=C["muted"])
        add_text(slide, cl + cw - Inches(1.9), y + Inches(0.2), Inches(1.4), Inches(0.25), deal, size=12, bold=True, color=C["pine"], align=PP_ALIGN.RIGHT)
        col = C["success"] if st == "Active" else (C["blue"] if st == "Scheduled" else C["muted"])
        add_text(slide, cl + cw - Inches(1.9), y + Inches(0.5), Inches(1.4), Inches(0.25), st, size=8, bold=True, color=col, align=PP_ALIGN.RIGHT)


def mock_payments(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Orders")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Payments · Stripe connected", size=11, bold=True, color=C["ink"])
    mock_kpi_row(
        slide,
        cl + Inches(0.15),
        ct + Inches(0.45),
        cw - Inches(0.3),
        [("Captured", "NZ$18.4k", "This week"), ("Pending", "NZ$620", "3 sessions"), ("Refunds", "NZ$84", "1 today"), ("Success rate", "99.2%", "30 days")],
    )
    add_rect(slide, cl + Inches(0.15), ct + Inches(1.45), cw - Inches(0.3), Inches(3.5), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.3), ct + Inches(1.6), Inches(4), Inches(0.25), "Stripe payout schedule", size=10, bold=True, color=C["ink"])
    add_text(slide, cl + Inches(0.3), ct + Inches(2.0), Inches(5), Inches(2.5), "• Card, Apple Pay, Google Pay at checkout\n• Server-side totals — never trust client prices\n• Automatic GST-ready receipts\n• Partial & full refunds from dashboard\n• Webhook-verified settlement\n• NZ bank payouts via Stripe", size=11, color=C["pine_deep"])


def mock_notifications(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Overview")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Notification centre", size=11, bold=True, color=C["ink"])
    notes = [
        ("Order #1842 confirmed", "Guest SMS + email", "Sent", C["success"]),
        ("Reservation reminder", "Wilson · tomorrow 7pm", "Queued", C["blue"]),
        ("Table ready alert", "Waitlist · Patel", "Sent", C["success"]),
        ("Refund processed", "Order #1831", "Sent", C["success"]),
        ("Staff alert", "86 Espresso Martini", "In-app", C["warn"]),
    ]
    for i, (title, sub, st, col) in enumerate(notes):
        y = ct + Inches(0.5) + i * Inches(0.85)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.75), C["white"], radius=0.08)
        add_rect(slide, cl + Inches(0.3), y + Inches(0.22), Inches(0.3), Inches(0.3), col, radius=0.5)
        add_text(slide, cl + Inches(0.8), y + Inches(0.12), Inches(3.5), Inches(0.25), title, size=11, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(0.8), y + Inches(0.4), Inches(3.5), Inches(0.22), sub, size=8, color=C["muted"])
        add_text(slide, cl + cw - Inches(1.3), y + Inches(0.28), Inches(0.9), Inches(0.25), st, size=9, bold=True, color=col)


def mock_reviews(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Analytics")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Reviews & reputation", size=11, bold=True, color=C["ink"])
    add_rect(slide, cl + Inches(0.15), ct + Inches(0.5), cw - Inches(0.3), Inches(1.2), C["bg_dark2"], radius=0.08)
    add_text(slide, cl + Inches(0.35), ct + Inches(0.7), Inches(2), Inches(0.4), "4.8", size=28, bold=True, color=C["gold"], font=FONT_DISPLAY)
    add_text(slide, cl + Inches(1.4), ct + Inches(0.85), Inches(3.5), Inches(0.6), "★★★★★  Google · 312 reviews\nResponse rate 94% · Avg reply 6h", size=10, color=C["muted_dark"])
    reviews = [
        ("“Best harbour dining in Auckland.”", "Ava C. · Google"),
        ("“QR ordering made Friday night effortless.”", "James P. · Email"),
        ("“Staff remembered our allergy notes.”", "Mia W. · Google"),
    ]
    for i, (q, who) in enumerate(reviews):
        y = ct + Inches(1.95) + i * Inches(0.95)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.85), C["white"], radius=0.08)
        add_text(slide, cl + Inches(0.3), y + Inches(0.15), cw - Inches(0.6), Inches(0.35), q, size=11, italic=True, color=C["ink"], font=FONT_DISPLAY)
        add_text(slide, cl + Inches(0.3), y + Inches(0.5), cw - Inches(0.6), Inches(0.25), who, size=8, color=C["muted"])


def mock_analytics(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Analytics")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Analytics & reports", size=11, bold=True, color=C["ink"])
    mock_kpi_row(
        slide,
        cl + Inches(0.15),
        ct + Inches(0.45),
        cw - Inches(0.3),
        [("Revenue", "NZ$42.8k", "+14% MoM"), ("AOV", "NZ$56.40", "+$3.20"), ("QR → order", "38%", "+5 pts"), ("Repeat guests", "41%", "+2 pts")],
    )
    # fake chart bars
    add_rect(slide, cl + Inches(0.15), ct + Inches(1.45), cw - Inches(0.3), Inches(2.3), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.3), ct + Inches(1.55), Inches(3), Inches(0.22), "Weekly revenue (NZ$)", size=9, bold=True, color=C["ink"])
    bars = [0.45, 0.62, 0.55, 0.78, 0.7, 0.92, 0.85]
    days = ["M", "T", "W", "T", "F", "S", "S"]
    max_h = Inches(1.5)
    for i, (b, d) in enumerate(zip(bars, days)):
        bh = int(max_h * b)
        x = cl + Inches(0.4) + i * Inches(0.7)
        y = ct + Inches(3.35) - bh
        add_rect(slide, x, y, Inches(0.45), bh, C["pine"] if i < 5 else C["gold"], radius=0.1)
        add_text(slide, x, ct + Inches(3.4), Inches(0.45), Inches(0.2), d, size=8, color=C["muted"], align=PP_ALIGN.CENTER)
    add_rect(slide, cl + Inches(0.15), ct + Inches(3.9), cw - Inches(0.3), Inches(1.1), C["white"], radius=0.08)
    add_text(slide, cl + Inches(0.3), ct + Inches(4.1), Inches(5), Inches(0.7), "Top items: Lamb Shoulder · Market Fish · Harbour Burger\nExport CSV · Date range filters · QR funnel included", size=10, color=C["pine_deep"])


def mock_staff(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Team")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Staff & roles", size=11, bold=True, color=C["ink"])
    roles = [
        ("Elena M.", "Owner", "Full access", C["gold"]),
        ("Tom K.", "Manager", "Ops + menu", C["pine"]),
        ("Priya S.", "Host", "Bookings only", C["blue"]),
        ("Jake R.", "Kitchen", "Orders + KDS", C["warn"]),
        ("Sam L.", "Front of house", "Orders view", C["muted"]),
    ]
    for i, (name, role, access, col) in enumerate(roles):
        y = ct + Inches(0.5) + i * Inches(0.85)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.75), C["white"], radius=0.08)
        add_rect(slide, cl + Inches(0.3), y + Inches(0.18), Inches(0.4), Inches(0.4), col, radius=0.5)
        add_text(slide, cl + Inches(0.35), y + Inches(0.25), Inches(0.35), Inches(0.28), name[0], size=11, bold=True, color=C["bg_dark"], align=PP_ALIGN.CENTER)
        add_text(slide, cl + Inches(0.9), y + Inches(0.12), Inches(2.5), Inches(0.25), name, size=11, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(0.9), y + Inches(0.4), Inches(2.5), Inches(0.25), role, size=9, color=C["muted"])
        add_text(slide, cl + cw - Inches(2.2), y + Inches(0.28), Inches(1.8), Inches(0.25), access, size=9, bold=True, color=C["pine"], align=PP_ALIGN.RIGHT)


def mock_inventory(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Menu")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Inventory & 86 board", size=11, bold=True, color=C["ink"])
    rows = [
        ("Lamb shoulder", "12 portions", "Low", C["warn"]),
        ("Market fish", "8 kg", "OK", C["success"]),
        ("Truffle oil", "2 bottles", "Critical", C["danger"]),
        ("House bread", "40 loaves", "OK", C["success"]),
        ("Espresso beans", "3 kg", "Low", C["warn"]),
        ("Oysters", "0 dozen", "86'd", C["danger"]),
    ]
    for i, (item, qty, st, col) in enumerate(rows):
        y = ct + Inches(0.5) + i * Inches(0.72)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.64), C["white"], radius=0.08)
        add_text(slide, cl + Inches(0.3), y + Inches(0.18), Inches(2.5), Inches(0.3), item, size=11, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(3.0), y + Inches(0.2), Inches(1.5), Inches(0.3), qty, size=10, color=C["muted"])
        add_rect(slide, cl + cw - Inches(1.5), y + Inches(0.16), Inches(1.1), Inches(0.32), col, radius=0.2)
        add_text(slide, cl + cw - Inches(1.5), y + Inches(0.18), Inches(1.1), Inches(0.28), st, size=8, bold=True, color=C["white"], align=PP_ALIGN.CENTER)


def mock_multiloc(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Overview")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Multi-location control", size=11, bold=True, color=C["ink"])
    locs = [
        ("Harbour Kitchen · Wynyard", "Auckland", "NZ$4.8k today", True),
        ("Harbour Kitchen · Britomart", "Auckland", "NZ$3.1k today", False),
        ("Harbour Kitchen · Queenstown", "Otago", "NZ$2.4k today", False),
    ]
    for i, (name, city, rev, on) in enumerate(locs):
        y = ct + Inches(0.55) + i * Inches(1.35)
        bg = C["bg_dark2"] if on else C["white"]
        fg = C["white"] if on else C["ink"]
        muted = C["muted_dark"] if on else C["muted"]
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(1.2), bg, radius=0.08)
        add_text(slide, cl + Inches(0.35), y + Inches(0.2), Inches(4), Inches(0.3), name, size=13, bold=True, color=C["gold"] if on else C["ink"])
        add_text(slide, cl + Inches(0.35), y + Inches(0.55), Inches(4), Inches(0.4), f"{city}  ·  {rev}", size=10, color=muted)
        add_text(slide, cl + cw - Inches(1.5), y + Inches(0.45), Inches(1.1), Inches(0.3), "Active" if on else "Switch", size=9, bold=True, color=C["gold"] if on else C["pine"], align=PP_ALIGN.RIGHT)


def mock_marketing(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Guests")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Marketing tools", size=11, bold=True, color=C["ink"])
    cards = [
        ("Email campaign", "Winter menu launch", "Open rate 42%"),
        ("SMS blast", "Tonight: last tables", "CTR 18%"),
        ("Segment", "Lapsed 60+ days", "126 guests"),
        ("Automation", "Post-visit thank you", "Active"),
    ]
    for i, (t1, t2, t3) in enumerate(cards):
        row, col = divmod(i, 2)
        x = cl + Inches(0.15) + col * ((cw - Inches(0.4)) / 2 + Inches(0.1))
        y = ct + Inches(0.55) + row * Inches(2.15)
        add_rect(slide, x, y, (cw - Inches(0.4)) / 2, Inches(1.95), C["white"], radius=0.08)
        add_rect(slide, x, y, (cw - Inches(0.4)) / 2, Inches(0.08), C["gold"])
        add_text(slide, x + Inches(0.2), y + Inches(0.3), Inches(2.5), Inches(0.25), t1, size=9, bold=True, color=C["pine"])
        add_text(slide, x + Inches(0.2), y + Inches(0.65), Inches(2.5), Inches(0.4), t2, size=14, bold=True, color=C["ink"], font=FONT_DISPLAY)
        add_text(slide, x + Inches(0.2), y + Inches(1.3), Inches(2.5), Inches(0.3), t3, size=10, color=C["muted"])


def mock_ai(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="AI")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "AI intelligence", size=11, bold=True, color=C["ink"])
    add_rect(slide, cl + Inches(0.15), ct + Inches(0.5), cw - Inches(0.3), Inches(1.6), C["bg_dark2"], radius=0.08)
    add_text(slide, cl + Inches(0.3), ct + Inches(0.65), Inches(4), Inches(0.25), "Tonight's briefing", size=10, bold=True, color=C["gold"])
    add_text(slide, cl + Inches(0.3), ct + Inches(1.0), Inches(5), Inches(0.9), "Friday covers are 18% above average. Lamb is on pace to 86 by 8:10pm. Suggest promoting Market Fish specials after 7:30.", size=11, color=C["muted_dark"])
    insights = [
        ("Menu copy draft", "Ready for owner approval"),
        ("Slow Tuesday forecast", "Promo window identified"),
        ("Guest segment tip", "Gold members prefer late seating"),
    ]
    for i, (t, s) in enumerate(insights):
        y = ct + Inches(2.35) + i * Inches(0.85)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.75), C["white"], radius=0.08)
        add_text(slide, cl + Inches(0.3), y + Inches(0.12), Inches(4), Inches(0.25), t, size=11, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(0.3), y + Inches(0.4), Inches(4), Inches(0.25), s, size=9, color=C["muted"])


def mock_security(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Team")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Security & permissions", size=11, bold=True, color=C["ink"])
    items = [
        ("Row-level security", "Every record scoped to your restaurant"),
        ("Role-based access", "Owner · Manager · Staff · Read-only"),
        ("Audit-ready logs", "Orders, refunds, menu publishes"),
        ("No raw card data", "Stripe handles PCI — never stored"),
        ("Server-side pricing", "Client prices are never trusted"),
        ("Secure auth", "Session-gated dashboard & APIs"),
    ]
    for i, (t, s) in enumerate(items):
        y = ct + Inches(0.5) + i * Inches(0.72)
        add_rect(slide, cl + Inches(0.15), y, cw - Inches(0.3), Inches(0.64), C["white"], radius=0.08)
        add_rect(slide, cl + Inches(0.3), y + Inches(0.17), Inches(0.3), Inches(0.3), C["pine"], radius=0.5)
        add_text(slide, cl + Inches(0.8), y + Inches(0.1), Inches(4), Inches(0.22), t, size=11, bold=True, color=C["ink"])
        add_text(slide, cl + Inches(0.8), y + Inches(0.35), Inches(4), Inches(0.22), s, size=8, color=C["muted"])


def mock_billing(slide, left, top, width, height):
    cl, ct, cw, ch = dash_shell(slide, left, top, width, height, active="Overview")
    add_text(slide, cl + Inches(0.15), ct + Inches(0.1), Inches(4), Inches(0.25), "Subscription & billing", size=11, bold=True, color=C["ink"])
    plans = [
        ("Starter", "NZ$149/mo", "1 location · core ops", False),
        ("Growth", "NZ$299/mo", "Ordering + CRM + AI", True),
        ("Group", "Custom", "Multi-site · SSO", False),
    ]
    for i, (n, price, desc, on) in enumerate(plans):
        x = cl + Inches(0.15) + i * ((cw - Inches(0.5)) / 3 + Inches(0.1))
        bg = C["bg_dark2"] if on else C["white"]
        add_rect(slide, x, ct + Inches(0.55), (cw - Inches(0.5)) / 3, Inches(4.3), bg, radius=0.08)
        add_text(slide, x + Inches(0.15), ct + Inches(0.8), Inches(1.6), Inches(0.3), n, size=12, bold=True, color=C["gold"] if on else C["ink"])
        add_text(slide, x + Inches(0.15), ct + Inches(1.3), Inches(1.6), Inches(0.4), price, size=16, bold=True, color=C["white"] if on else C["ink"], font=FONT_DISPLAY)
        add_text(slide, x + Inches(0.15), ct + Inches(1.9), Inches(1.6), Inches(0.8), desc, size=10, color=C["muted_dark"] if on else C["muted"])
        if on:
            add_rect(slide, x + Inches(0.15), ct + Inches(4.3), Inches(1.5), Inches(0.32), C["gold"], radius=0.2)
            add_text(slide, x + Inches(0.15), ct + Inches(4.32), Inches(1.5), Inches(0.28), "Current plan", size=8, bold=True, color=C["bg_dark"], align=PP_ALIGN.CENTER)


def mock_mobile(slide, left, top, width, height):
    # phone frame
    phone_w = Inches(2.4)
    phone_h = Inches(5.1)
    px = left + (width - phone_w) / 2
    py = top + Inches(0.35)
    add_rect(slide, px, py, phone_w, phone_h, C["bg_dark"], radius=0.12, shadow=True)
    add_rect(slide, px + Inches(0.08), py + Inches(0.08), phone_w - Inches(0.16), phone_h - Inches(0.16), C["cream"])
    add_rect(slide, px + Inches(0.7), py + Inches(0.18), Inches(1.0), Inches(0.12), C["bg_dark"], radius=0.5)
    add_text(slide, px + Inches(0.2), py + Inches(0.45), phone_w - Inches(0.4), Inches(0.25), "Harbour Kitchen", size=10, bold=True, color=C["ink"], align=PP_ALIGN.CENTER)
    add_text(slide, px + Inches(0.2), py + Inches(0.75), phone_w - Inches(0.4), Inches(0.4), "Tonight's service", size=14, bold=True, color=C["ink"], font=FONT_DISPLAY, align=PP_ALIGN.CENTER)
    for i, (lab, val) in enumerate([("Open tickets", "12"), ("Ready pickup", "3"), ("Waitlist", "7")]):
        y = py + Inches(1.35) + i * Inches(0.85)
        add_rect(slide, px + Inches(0.2), y, phone_w - Inches(0.4), Inches(0.72), C["white"], radius=0.1)
        add_text(slide, px + Inches(0.35), y + Inches(0.12), Inches(1.5), Inches(0.22), lab, size=8, color=C["muted"])
        add_text(slide, px + Inches(0.35), y + Inches(0.35), Inches(1.5), Inches(0.28), val, size=16, bold=True, color=C["ink"])
    add_rect(slide, px + Inches(0.2), py + Inches(4.1), phone_w - Inches(0.4), Inches(0.55), C["pine"], radius=0.15)
    add_text(slide, px + Inches(0.2), py + Inches(4.22), phone_w - Inches(0.4), Inches(0.35), "Open kitchen board", size=9, bold=True, color=C["white"], align=PP_ALIGN.CENTER)
    # tablet hint
    add_text(slide, left, top + height - Inches(0.45), width, Inches(0.3), "Optimised for iPad POS / host stand · phone for managers on the floor", size=9, color=C["muted"], align=PP_ALIGN.CENTER)


# ── Slide builders ──────────────────────────────────────────────────────────

def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["bg_dark"])
    # aurora-ish accents
    add_rect(s, Inches(8.5), Inches(-1), Inches(6), Inches(4), C["bg_dark2"], radius=0.3)
    add_rect(s, Inches(-1), Inches(5), Inches(5), Inches(4), C["bg_panel"], radius=0.3)
    accent_bar(s)
    add_text(s, MARGIN, Inches(1.6), Inches(8), Inches(0.35), "BUILT FOR NEW ZEALAND HOSPITALITY", size=11, bold=True, color=C["gold"])
    add_text(s, MARGIN, Inches(2.15), Inches(9), Inches(1.1), "Restaurant OS", size=54, bold=True, color=C["white"], font=FONT_DISPLAY)
    add_text(
        s,
        MARGIN,
        Inches(3.4),
        Inches(8),
        Inches(0.9),
        "The all-in-one operating system for modern restaurants —\nwebsite, menu, orders, bookings, guests, and growth.",
        size=18,
        color=C["muted_dark"],
    )
    add_rect(s, MARGIN, Inches(4.7), Inches(2.4), Inches(0.48), C["gold"], radius=0.15)
    add_text(s, MARGIN, Inches(4.8), Inches(2.4), Inches(0.35), "Client presentation", size=12, bold=True, color=C["bg_dark"], align=PP_ALIGN.CENTER)
    add_text(s, MARGIN, Inches(6.6), Inches(8), Inches(0.3), "Sample venue: Harbour Kitchen · Wynyard Quarter, Auckland", size=11, color=C["muted"])
    # right product preview mini
    mock_dashboard(s, Inches(7.0), Inches(1.5), Inches(5.7), Inches(4.8))


def slide_pain_points(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=False)
    eyebrow(s, MARGIN, Inches(0.85), "The problem", dark=False)
    section_title(s, MARGIN, Inches(1.15), "What restaurant owners struggle with", dark=False, width=Inches(11))
    pains = [
        ("Too many tools", "Website builder, menu PDF, booking widget, delivery app, POS notes — nothing talks."),
        ("Menu chaos", "Printed menus and third-party apps go stale the moment a dish is 86'd."),
        ("Missed revenue", "No online ordering, weak upsell, and guests bounce to aggregators."),
        ("Blind service", "No single view of covers, tickets, waitlist, or guest history."),
        ("Manual grind", "Staff re-type bookings, chase confirmations, and answer the same phone calls."),
        ("No clear ROI", "Hard to see what drives spend, repeat visits, or quiet nights."),
    ]
    for i, (t, b) in enumerate(pains):
        row, col = divmod(i, 3)
        x = MARGIN + col * Inches(4.1)
        y = Inches(2.2) + row * Inches(2.3)
        card(s, x, y, Inches(3.9), Inches(2.05))
        add_rect(s, x, y, Inches(0.1), Inches(2.05), C["gold"])
        add_text(s, x + Inches(0.3), y + Inches(0.3), Inches(3.3), Inches(0.4), t, size=16, bold=True, color=C["ink"], font=FONT_DISPLAY)
        add_text(s, x + Inches(0.3), y + Inches(0.85), Inches(3.3), Inches(0.9), b, size=12, color=C["pine_deep"])


def slide_fragmented(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["bg_dark"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=True)
    eyebrow(s, MARGIN, Inches(0.85), "Today's reality")
    section_title(s, MARGIN, Inches(1.15), "A fragmented workflow tax")
    body(s, MARGIN, Inches(1.9), Inches(11), "Most venues run 6–10 disconnected products. Every handoff costs time, money, and guest experience.")
    tools = ["Website", "Menu PDF", "OpenTable", "Uber Eats", "WhatsApp", "Spreadsheets", "POS", "Mailchimp"]
    for i, t in enumerate(tools):
        x = MARGIN + (i % 4) * Inches(3.05)
        y = Inches(2.9) + (i // 4) * Inches(1.5)
        add_rect(s, x, y, Inches(2.85), Inches(1.2), C["bg_panel"], radius=0.1)
        add_text(s, x, y + Inches(0.4), Inches(2.85), Inches(0.4), t, size=16, bold=True, color=C["white"], align=PP_ALIGN.CENTER, font=FONT_DISPLAY)
    add_text(s, MARGIN, Inches(6.4), Inches(11), Inches(0.4), "Result: inconsistent brand · double entry · delayed updates · aggregator fees · no single source of truth", size=13, color=C["gold"])


def slide_solution(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=False)
    eyebrow(s, MARGIN, Inches(0.85), "The solution", dark=False)
    section_title(s, MARGIN, Inches(1.15), "One platform. One source of truth.", dark=False)
    body(s, MARGIN, Inches(1.9), Inches(7), "Restaurant OS replaces the tool stack with a single operating system — guest-facing and staff-facing — built for NZ operators.", dark=False)
    pillars = [
        ("Attract", "Premium website, SEO, QR menus, offers"),
        ("Convert", "Ordering, reservations, waitlist, payments"),
        ("Operate", "Kitchen, staff roles, inventory, multi-site"),
        ("Grow", "CRM, loyalty, analytics, AI insights"),
    ]
    for i, (t, b) in enumerate(pillars):
        x = MARGIN + i * Inches(3.1)
        card(s, x, Inches(3.5), Inches(2.95), Inches(2.6))
        add_text(s, x + Inches(0.2), Inches(3.75), Inches(2.5), Inches(0.2), f"0{i+1}", size=11, bold=True, color=C["gold"])
        add_text(s, x + Inches(0.2), Inches(4.15), Inches(2.5), Inches(0.4), t, size=20, bold=True, color=C["ink"], font=FONT_DISPLAY)
        add_text(s, x + Inches(0.2), Inches(4.75), Inches(2.5), Inches(0.9), b, size=13, color=C["pine_deep"])


def slide_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["bg_dark"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=True)
    eyebrow(s, MARGIN, Inches(0.85), "Platform")
    section_title(s, MARGIN, Inches(1.15), "Restaurant OS at a glance")
    groups = [
        ("Guest experience", ["Branded website", "Digital & QR menu", "Online ordering", "Reservations"]),
        ("Service ops", ["Order & KDS", "Pickup / delivery", "Waitlist", "Staff roles"]),
        ("Growth", ["CRM & loyalty", "Offers", "Reviews", "Marketing"]),
        ("Intelligence", ["Analytics", "AI insights", "Multi-location", "Secure billing"]),
    ]
    for i, (title, items) in enumerate(groups):
        x = MARGIN + i * Inches(3.1)
        add_rect(s, x, Inches(2.2), Inches(2.95), Inches(4.4), C["bg_panel"], radius=0.08)
        add_text(s, x + Inches(0.2), Inches(2.45), Inches(2.5), Inches(0.4), title, size=14, bold=True, color=C["gold"])
        for j, it in enumerate(items):
            add_text(s, x + Inches(0.2), Inches(3.1) + j * Inches(0.7), Inches(2.5), Inches(0.5), f"▸  {it}", size=13, color=C["white"])


def slide_ecosystem(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=False)
    eyebrow(s, MARGIN, Inches(0.85), "Ecosystem", dark=False)
    section_title(s, MARGIN, Inches(1.15), "Everything connected", dark=False)
    # center hub
    cx, cy = Inches(5.9), Inches(4.2)
    add_rect(s, cx, cy, Inches(2.4), Inches(1.1), C["bg_dark"], radius=0.12)
    add_text(s, cx, cy + Inches(0.3), Inches(2.4), Inches(0.5), "Restaurant OS", size=14, bold=True, color=C["gold"], align=PP_ALIGN.CENTER, font=FONT_DISPLAY)
    nodes = [
        (Inches(0.7), Inches(2.3), "Website"),
        (Inches(3.2), Inches(2.3), "Menu"),
        (Inches(8.5), Inches(2.3), "Orders"),
        (Inches(10.8), Inches(2.3), "Payments"),
        (Inches(0.7), Inches(5.5), "CRM"),
        (Inches(3.2), Inches(5.5), "Loyalty"),
        (Inches(8.5), Inches(5.5), "Analytics"),
        (Inches(10.8), Inches(5.5), "AI"),
        (Inches(5.5), Inches(2.0), "QR"),
        (Inches(5.5), Inches(5.9), "Bookings"),
    ]
    for x, y, label in nodes:
        add_rect(s, x, y, Inches(2.0), Inches(0.7), C["white"], line=C["line"], radius=0.12, shadow=True)
        add_text(s, x, y + Inches(0.18), Inches(2.0), Inches(0.4), label, size=13, bold=True, color=C["ink"], align=PP_ALIGN.CENTER)


def slide_workflow(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["bg_dark"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=True)
    eyebrow(s, MARGIN, Inches(0.85), "End-to-end")
    section_title(s, MARGIN, Inches(1.15), "From discovery to loyalty")
    steps = [
        ("01", "Discover", "Website / QR / search"),
        ("02", "Browse", "Live digital menu"),
        ("03", "Book / Order", "Table or takeaway"),
        ("04", "Pay", "Stripe checkout"),
        ("05", "Serve", "KDS & floor tools"),
        ("06", "Retain", "CRM · loyalty · AI"),
    ]
    for i, (n, t, b) in enumerate(steps):
        x = MARGIN + i * Inches(2.1)
        add_rect(s, x, Inches(2.6), Inches(1.95), Inches(3.2), C["bg_panel"], radius=0.1)
        add_text(s, x + Inches(0.15), Inches(2.85), Inches(1.6), Inches(0.3), n, size=12, bold=True, color=C["gold"])
        add_text(s, x + Inches(0.15), Inches(3.4), Inches(1.65), Inches(0.6), t, size=16, bold=True, color=C["white"], font=FONT_DISPLAY)
        add_text(s, x + Inches(0.15), Inches(4.3), Inches(1.65), Inches(1.0), b, size=12, color=C["muted_dark"])
        if i < len(steps) - 1:
            add_text(s, x + Inches(1.75), Inches(3.8), Inches(0.4), Inches(0.3), "→", size=16, color=C["gold"])


def slide_before_after(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=False)
    eyebrow(s, MARGIN, Inches(0.85), "Transformation", dark=False)
    section_title(s, MARGIN, Inches(1.15), "Before vs After Restaurant OS", dark=False)
    # before
    add_rect(s, MARGIN, Inches(2.2), Inches(5.8), Inches(4.5), C["white"], line=C["line"], radius=0.08, shadow=True)
    add_text(s, MARGIN + Inches(0.35), Inches(2.45), Inches(5), Inches(0.4), "Before", size=20, bold=True, color=C["danger"], font=FONT_DISPLAY)
    before = [
        "5–10 logins every service",
        "Stale menus & broken booking links",
        "Aggregator fees eating margin",
        "Guest data trapped in inboxes",
        "Managers flying blind on covers",
        "Marketing is an afterthought",
    ]
    for i, t in enumerate(before):
        add_text(s, MARGIN + Inches(0.35), Inches(3.1) + i * Inches(0.5), Inches(5.2), Inches(0.4), f"✕  {t}", size=13, color=C["ink"])
    # after
    add_rect(s, Inches(6.9), Inches(2.2), Inches(5.8), Inches(4.5), C["bg_dark"], radius=0.08, shadow=True)
    add_text(s, Inches(7.25), Inches(2.45), Inches(5), Inches(0.4), "After", size=20, bold=True, color=C["gold"], font=FONT_DISPLAY)
    after = [
        "One dashboard for the whole venue",
        "Live menu, QR, and site in sync",
        "Direct orders with Stripe payouts",
        "Unified guest profiles & loyalty",
        "Live boards for kitchen & floor",
        "AI + analytics guiding decisions",
    ]
    for i, t in enumerate(after):
        add_text(s, Inches(7.25), Inches(3.1) + i * Inches(0.5), Inches(5.2), Inches(0.4), f"✓  {t}", size=13, color=C["white"])


def slide_benefits(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["bg_dark"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=True)
    eyebrow(s, MARGIN, Inches(0.85), "Value")
    section_title(s, MARGIN, Inches(1.15), "Key business benefits")
    items = [
        ("↑ Direct revenue", "Own the guest relationship — less platform tax."),
        ("↓ Operational load", "Fewer tools, fewer errors, faster service."),
        ("↑ Table yield", "Smarter bookings, waitlist, and turn times."),
        ("↑ Repeat visits", "CRM, loyalty, and timely offers that feel personal."),
        ("Clarity", "Live analytics and AI briefings for owners."),
        ("Brand control", "Premium digital presence that matches the room."),
    ]
    for i, (t, b) in enumerate(items):
        row, col = divmod(i, 3)
        x = MARGIN + col * Inches(4.1)
        y = Inches(2.2) + row * Inches(2.3)
        add_rect(s, x, y, Inches(3.9), Inches(2.05), C["bg_panel"], radius=0.1)
        add_text(s, x + Inches(0.3), y + Inches(0.35), Inches(3.3), Inches(0.45), t, size=18, bold=True, color=C["gold"], font=FONT_DISPLAY)
        add_text(s, x + Inches(0.3), y + Inches(1.0), Inches(3.3), Inches(0.7), b, size=13, color=C["muted_dark"])


def slide_why(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=False)
    eyebrow(s, MARGIN, Inches(0.85), "Why us", dark=False)
    section_title(s, MARGIN, Inches(1.15), "Why Restaurant OS", dark=False)
    reasons = [
        ("NZ-first", "Timezone, GST mindset, Auckland-to-Queenstown multi-site reality."),
        ("Premium guest UX", "Your brand — not a generic booking widget."),
        ("Ops-grade core", "Server-side pricing, Stripe-ready payments, role security."),
        ("Owner-approved AI", "Insights and drafts that never publish without you."),
        ("One contract", "Website + menu + orders + bookings + CRM — not five vendors."),
        ("Built to scale", "From a single bistro to a regional group."),
    ]
    for i, (t, b) in enumerate(reasons):
        row, col = divmod(i, 2)
        x = MARGIN + col * Inches(6.2)
        y = Inches(2.15) + row * Inches(1.55)
        card(s, x, y, Inches(5.95), Inches(1.4))
        add_text(s, x + Inches(0.3), y + Inches(0.25), Inches(5.3), Inches(0.35), t, size=16, bold=True, color=C["ink"], font=FONT_DISPLAY)
        add_text(s, x + Inches(0.3), y + Inches(0.7), Inches(5.3), Inches(0.5), b, size=13, color=C["pine_deep"])


def slide_roadmap(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["bg_dark"])
    accent_bar(s)
    brand_chip(s, MARGIN, Inches(0.28), dark=True)
    eyebrow(s, MARGIN, Inches(0.85), "Roadmap")
    section_title(s, MARGIN, Inches(1.15), "Where we're headed")
    phases = [
        ("Now", "Website · Menu · QR · Orders · Reservations · CRM · Analytics · AI briefs", C["gold"]),
        ("Next", "Live Stripe · SMS/email providers · Waitlist polish · Loyalty · Offers engine", C["pine"]),
        ("Later", "Full KDS packs · Inventory depth · Multi-brand groups · Advanced marketing automation", C["muted"]),
    ]
    for i, (title, body, col) in enumerate(phases):
        y = Inches(2.3) + i * Inches(1.45)
        add_rect(s, MARGIN, y, Inches(12.2), Inches(1.25), C["bg_panel"], radius=0.1)
        add_rect(s, MARGIN, y, Inches(0.12), Inches(1.25), col)
        add_text(s, MARGIN + Inches(0.4), y + Inches(0.25), Inches(2), Inches(0.35), title, size=18, bold=True, color=col, font=FONT_DISPLAY)
        add_text(s, MARGIN + Inches(2.5), y + Inches(0.35), Inches(9.2), Inches(0.6), body, size=14, color=C["white"])


def slide_cta(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["bg_dark"])
    accent_bar(s)
    add_rect(s, Inches(7), Inches(-0.5), Inches(8), Inches(4), C["bg_panel"], radius=0.3)
    add_text(s, MARGIN, Inches(2.0), Inches(10), Inches(0.4), "LET'S GET YOUR VENUE LIVE", size=12, bold=True, color=C["gold"])
    add_text(s, MARGIN, Inches(2.55), Inches(10), Inches(1.0), "Ready to run your restaurant\non one operating system?", size=34, bold=True, color=C["white"], font=FONT_DISPLAY)
    add_text(s, MARGIN, Inches(4.0), Inches(9), Inches(0.7), "Book a walkthrough with Harbour Kitchen demo data —\nor start your own workspace.", size=16, color=C["muted_dark"])
    add_rect(s, MARGIN, Inches(5.1), Inches(2.8), Inches(0.55), C["gold"], radius=0.15)
    add_text(s, MARGIN, Inches(5.2), Inches(2.8), Inches(0.4), "hello@restaurantos.nz", size=13, bold=True, color=C["bg_dark"], align=PP_ALIGN.CENTER)
    add_text(s, MARGIN + Inches(3.1), Inches(5.25), Inches(5), Inches(0.4), "Auckland · New Zealand", size=14, color=C["muted_dark"])
    add_text(s, MARGIN, Inches(6.5), Inches(10), Inches(0.3), "Restaurant OS  ·  Premium hospitality software for Aotearoa", size=11, color=C["muted"])


FEATURE_SLIDES = [
    ("Owner / Admin Dashboard", "Command centre", ["See today's revenue, orders, covers, and QR activity in one view.", "Jump straight into live tickets, bookings, and AI briefings.", "Run the floor without hopping between apps."], "No single source of truth during service", "Faster decisions, fewer missed tickets", mock_dashboard),
    ("Restaurant Profile & Settings", "Brand & venue setup", ["Manage logo, colours, hours, timezone, and contact details.", "Keep GST, currency, and location settings accurate.", "Publish changes that flow to the guest website instantly."], "Brand inconsistency across channels", "One accurate venue profile everywhere", mock_settings),
    ("Website Builder", "Guest-facing brand", ["Launch a premium branded site — home, menu, gallery, contact.", "SEO metadata, sitemap, and custom-domain ready architecture.", "Preview unpublished changes before go-live."], "Generic templates that don't match the room", "A digital front door that converts", mock_website),
    ("Digital Menu Management", "Live menu control", ["Create menus, categories, items, modifiers, and photos.", "Drag-and-drop reorder · 86 items in seconds.", "Publish once — website, QR, and ordering stay in sync."], "Stale PDFs and third-party menu lag", "Always-accurate menus, less waste", mock_menu),
    ("QR Menu", "Table-side ordering", ["Generate unique QR codes per table.", "Guests open the live menu with table context.", "Track scans and convert to dine-in orders."], "Printed menus and slow service", "Higher turn speed, contactless convenience", mock_qr),
    ("Online Ordering", "Direct digital sales", ["Pickup, delivery, and dine-in flows on your brand.", "Server-side pricing with GST-aware totals.", "Secure checkout without sending guests to aggregators."], "Aggregator fees and brand leakage", "Higher margin direct orders", mock_ordering),
    ("Pickup & Delivery Management", "Fulfilment board", ["Separate pickup and delivery queues in real time.", "Assign status, ETAs, and driver notes.", "Keep the pass calm when volume spikes."], "Chaos between kitchen and runners", "On-time fulfilments, happier guests", mock_fulfillment),
    ("Order Management / Kitchen Display", "Kitchen command", ["Ticket cards with elapsed time and station clarity.", "Bump completed items · prioritise aging tickets.", "Works on kitchen screens and manager tablets."], "Paper tickets and missed mods", "Faster ticket times, fewer errors", mock_kds),
    ("Reservations", "Bookings that fit the floor", ["Online booking with covers, hours, and slot rules.", "Confirm, seat, reschedule, or cancel from the dashboard.", "Deposit-ready private dining support."], "Double-bookings and phone tag", "Fuller books with controlled capacity", mock_reservations),
    ("Waitlist Management", "When you're full", ["Capture walk-ins with party size and quotes.", "SMS when the table is ready.", "Keep the door calm on peak nights."], "Crowded entrances and lost walk-ins", "Recover covers you'd otherwise turn away", mock_waitlist),
    ("Customer / CRM", "Know your guests", ["Auto-built profiles from orders and reservations.", "See spend, visits, allergies, and preferences.", "Serve regulars like VIPs — every time."], "Guest history scattered across tools", "Personal service that drives loyalty", mock_crm),
    ("Customer Loyalty", "Reasons to return", ["Points, tiers, and redeemable rewards.", "Birthday and VIP recognition built in.", "Visible on guest profiles and receipts."], "One-off visits with no reason to return", "Higher repeat rate and LTV", mock_loyalty),
    ("Offers & Promotions", "Fill quiet sessions", ["Launch lunch sets, happy hours, and seasonal bundles.", "Schedule campaigns · target loyalty segments.", "Measure redemption against revenue."], "Empty midweek covers", "Demand shaping without discount chaos", mock_offers),
    ("Payments & Stripe", "Money, handled properly", ["Card and wallet payments via Stripe.", "Refunds, payouts, and webhook-verified settlement.", "No raw card data stored in Restaurant OS."], "Untrusted checkout and reconciliation pain", "Reliable cashflow with PCI burden offloaded", mock_payments),
    ("Notifications", "The right ping", ["Order and booking updates to guests and staff.", "SMS / email templates with retry-safe delivery.", "In-app alerts for 86s and service events."], "Missed confirmations and no-shows", "Fewer no-shows, clearer communication", mock_notifications),
    ("Reviews & Reputation", "Protect the brand", ["Monitor ratings and recent guest feedback.", "Respond quickly from one inbox.", "Spot service themes before they trend."], "Silent reviews damaging discovery", "Stronger public reputation", mock_reviews),
    ("Analytics & Reports", "See what works", ["Revenue, AOV, item performance, QR funnel.", "Date ranges and CSV export for accountants.", "Facts — never invented metrics."], "Gut-feel decisions", "Evidence-based ops and menu design", mock_analytics),
    ("Staff & Roles", "Right access, right people", ["Invite owners, managers, hosts, and kitchen staff.", "Limit permissions by role.", "Keep sensitive actions owner-only."], "Shared passwords and over-access", "Safer ops with clear accountability", mock_staff),
    ("Inventory Management", "Stop 86 surprises", ["Track critical stock and portion counts.", "Flag low / critical / 86'd items.", "Sync availability back to the live menu."], "Selling dishes you can't plate", "Less waste, fewer embarrassed apologies", mock_inventory),
    ("Multi-location Management", "Group control", ["Switch between venues without losing context.", "Compare today's performance across sites.", "Shared brand standards, local hours & menus."], "Spreadsheets to run a group", "Consistent ops at regional scale", mock_multiloc),
    ("Marketing Tools", "Own your demand", ["Email and SMS campaigns on your guest list.", "Segments for lapsed, VIP, and local guests.", "Automations like post-visit thank-yous."], "Paying platforms to reach your own guests", "Lower CAC, higher owned demand", mock_marketing),
    ("AI-Powered Features", "A quiet second brain", ["Service briefings from real venue data.", "Menu copy drafts for owner approval.", "Forecast quiet sessions and promo windows.", "AI never publishes without you."], "No time to analyse the numbers", "Better decisions without hiring an analyst", mock_ai),
    ("Security & Permissions", "Enterprise-grade trust", ["Tenant isolation and membership-based access.", "Server-side pricing and payment verification.", "Audit-friendly operational history."], "Fear of leaks and chargeback risk", "Confidence to run digitally", mock_security),
    ("Subscription & Billing", "Simple commercial terms", ["Clear plans for single sites and groups.", "Upgrade as you add ordering, CRM, and AI.", "Transparent NZD pricing."], "Surprise vendor fees", "Predictable software cost", mock_billing),
    ("Mobile / Tablet Experience", "Service in your pocket", ["Manager phone view for live service metrics.", "Tablet-ready kitchen and host boards.", "Responsive guest ordering on any device."], "Desktop-only tools that fail on the floor", "Ops visibility wherever you stand", mock_mobile),
]


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_cover(prs)
    slide_pain_points(prs)
    slide_fragmented(prs)
    slide_solution(prs)
    slide_overview(prs)
    slide_ecosystem(prs)

    for title, eye, bullets, pain, benefit, mock in FEATURE_SLIDES:
        # alternate light/dark lightly — keep light for readability of mockups
        feature_slide(
            prs,
            title=title,
            eyebrow_text=eye,
            bullets=bullets,
            pain=pain,
            benefit=benefit,
            mock_fn=mock,
            dark=False,
        )

    slide_workflow(prs)
    slide_before_after(prs)
    slide_benefits(prs)
    slide_why(prs)
    slide_roadmap(prs)
    slide_cta(prs)

    out = Path(__file__).resolve().parents[1] / "docs" / "presentations" / "Restaurant-OS-Platform-Presentation.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"Wrote {out} ({len(prs.slides)} slides)")
    return out


if __name__ == "__main__":
    build()
