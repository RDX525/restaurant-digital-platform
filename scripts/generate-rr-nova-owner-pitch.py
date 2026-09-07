#!/usr/bin/env python3
"""Generate a premium 10-slide RR NOVA sales deck for restaurant owners."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

# ── Design tokens — warm cream · charcoal · muted gold ──────────────────────
C = {
    "cream": RGBColor(0xF6, 0xF1, 0xE8),
    "cream_deep": RGBColor(0xEB, 0xE3, 0xD6),
    "cream_soft": RGBColor(0xFB, 0xF8, 0xF3),
    "charcoal": RGBColor(0x1C, 0x19, 0x16),
    "charcoal_soft": RGBColor(0x2E, 0x2A, 0x25),
    "ink": RGBColor(0x2A, 0x26, 0x21),
    "muted": RGBColor(0x6E, 0x66, 0x5A),
    "muted_light": RGBColor(0x9A, 0x91, 0x84),
    "gold": RGBColor(0xB8, 0x97, 0x64),
    "gold_soft": RGBColor(0xD4, 0xC0, 0x98),
    "gold_deep": RGBColor(0x9A, 0x7B, 0x4A),
    "line": RGBColor(0xE0, 0xD6, 0xC6),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "danger_soft": RGBColor(0x8C, 0x5A, 0x4A),
}

FONT_DISPLAY = "Georgia"
FONT_SANS = "Avenir Next"

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "presentations" / "assets"
OUT_PPTX = ROOT / "docs" / "presentations" / "RR-NOVA-Restaurant-Owner-Pitch.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.7)
DEMO_URL = "restaurant-digital-platform-beige.vercel.app"


def _set_run_font(run, size, *, bold=False, color=None, font=FONT_SANS, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color


def _fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _line(shape, color: RGBColor | None = None, width_pt: float = 0.75):
    if color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color
        shape.line.width = Pt(width_pt)


def _set_shape_shadow(shape):
    spPr = shape._element.spPr
    effectLst = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    if effectLst is None:
        effectLst = OxmlElement("a:effectLst")
        spPr.append(effectLst)
    for child in list(effectLst):
        effectLst.remove(child)
    outer = OxmlElement("a:outerShdw")
    outer.set("blurRad", "63500")
    outer.set("dist", "38100")
    outer.set("dir", "2700000")
    outer.set("algn", "tl")
    outer.set("rotWithShape", "0")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", "1C1916")
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "22000")
    srgb.append(alpha)
    outer.append(srgb)
    effectLst.append(outer)


def add_rect(slide, left, top, width, height, fill, line=None, radius=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    _fill(shp, fill)
    _line(shp, line)
    if radius is not None and hasattr(shp, "adjustments") and len(shp.adjustments) > 0:
        shp.adjustments[0] = radius
    if shadow:
        _set_shape_shadow(shp)
    return shp


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=14,
    bold=False,
    color=C["ink"],
    font=FONT_SANS,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    italic=False,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, bold=bold, color=color, font=font, italic=italic)
    return box


def add_paras(slide, left, top, width, height, lines: list[tuple[str, dict]]):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, style) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = style.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(style.get("space_before", 0))
        p.space_after = Pt(style.get("space_after", 6))
        run = p.add_run()
        run.text = text
        _set_run_font(
            run,
            style.get("size", 14),
            bold=style.get("bold", False),
            color=style.get("color", C["ink"]),
            font=style.get("font", FONT_SANS),
            italic=style.get("italic", False),
        )
    return box


def solid_bg(slide, color):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, color)


def gold_rule(slide, left, top, width=Inches(1.2)):
    add_rect(slide, left, top, width, Pt(2), C["gold"])


def brand_mark(slide, left, top, *, dark=False):
    """Minimal RR NOVA wordmark."""
    add_rect(slide, left, top + Inches(0.06), Inches(0.22), Inches(0.22), C["gold"], radius=0.15)
    add_text(
        slide,
        left + Inches(0.34),
        top,
        Inches(2.5),
        Inches(0.34),
        "RR NOVA",
        size=16,
        bold=True,
        color=C["cream"] if dark else C["charcoal"],
        font=FONT_DISPLAY,
    )


def slide_number(slide, n, total=10, *, dark=False):
    add_text(
        slide,
        Inches(11.6),
        Inches(7.05),
        Inches(1.2),
        Inches(0.3),
        f"{n:02d}  /  {total:02d}",
        size=10,
        color=C["muted_light"] if not dark else C["muted"],
        align=PP_ALIGN.RIGHT,
    )


def picture_frame(slide, path: Path, left, top, width, height):
    """Inset product screenshot with subtle frame."""
    frame = add_rect(
        slide,
        left - Inches(0.06),
        top - Inches(0.06),
        width + Inches(0.12),
        height + Inches(0.12),
        C["white"],
        line=C["line"],
        radius=0.04,
        shadow=True,
    )
    pic = slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    return frame, pic


# ── Slides ───────────────────────────────────────────────────────────────────


def slide_01_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["charcoal"])

    # Soft cream panel on the left for typography calm
    add_rect(s, 0, 0, Inches(6.6), SLIDE_H, C["cream"])

    # Gold accent line
    add_rect(s, 0, 0, Inches(0.08), SLIDE_H, C["gold"])

    brand_mark(s, MARGIN, Inches(0.55), dark=False)

    add_text(
        s,
        MARGIN,
        Inches(1.9),
        Inches(5.4),
        Inches(0.35),
        "FOR RESTAURANT OWNERS",
        size=11,
        bold=True,
        color=C["gold_deep"],
    )

    add_paras(
        s,
        MARGIN,
        Inches(2.4),
        Inches(5.5),
        Inches(2.4),
        [
            (
                "Your restaurant.",
                {"size": 40, "bold": True, "color": C["charcoal"], "font": FONT_DISPLAY, "space_after": 2},
            ),
            (
                "Your customers.",
                {"size": 40, "bold": True, "color": C["charcoal"], "font": FONT_DISPLAY, "space_after": 2},
            ),
            (
                "Your data.",
                {"size": 40, "bold": True, "color": C["charcoal"], "font": FONT_DISPLAY, "space_after": 14},
            ),
            (
                "One platform.",
                {"size": 22, "bold": True, "color": C["gold_deep"], "font": FONT_DISPLAY, "space_after": 0},
            ),
        ],
    )

    add_text(
        s,
        MARGIN,
        Inches(5.85),
        Inches(5.4),
        Inches(0.7),
        "The all-in-one platform that turns your website,\nmenu, orders, and bookings into one growth engine.",
        size=14,
        color=C["muted"],
    )

    add_text(
        s,
        MARGIN,
        Inches(6.85),
        Inches(5.4),
        Inches(0.3),
        "Crafted for Aotearoa New Zealand hospitality",
        size=11,
        color=C["muted_light"],
        italic=True,
    )

    # Right visual — full-bleed hospitality image
    hero = ASSETS / "02-harbour-kitchen.png"
    if hero.exists():
        s.shapes.add_picture(str(hero), Inches(6.6), Inches(0), height=SLIDE_H)

    # Dark gradient-feel strip for legibility of footer
    add_rect(s, Inches(6.6), Inches(6.55), Inches(6.8), Inches(0.95), C["charcoal"])
    add_text(
        s,
        Inches(7.0),
        Inches(6.85),
        Inches(5.5),
        Inches(0.3),
        "Live demo · Harbour Kitchen · Auckland",
        size=11,
        color=C["gold_soft"],
    )
    slide_number(s, 1, dark=True)


def slide_02_problem(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.4))
    slide_number(s, 2)

    add_text(s, MARGIN, Inches(1.15), Inches(8), Inches(0.3), "THE PROBLEM", size=11, bold=True, color=C["gold_deep"])
    gold_rule(s, MARGIN, Inches(1.5))
    add_text(
        s,
        MARGIN,
        Inches(1.7),
        Inches(11),
        Inches(0.7),
        "Too many disconnected systems.",
        size=36,
        bold=True,
        color=C["charcoal"],
        font=FONT_DISPLAY,
    )
    add_text(
        s,
        MARGIN,
        Inches(2.45),
        Inches(10),
        Inches(0.4),
        "Website builders, PDF menus, booking widgets, delivery apps, and spreadsheets — nothing talks to each other.",
        size=15,
        color=C["muted"],
    )

    pains = [
        ("Too many tools", "5–10 logins every service — nothing talks to each other."),
        ("Menu chaos", "Printed menus and apps go stale the moment a dish is 86’d."),
        ("Missed revenue", "Guests bounce to aggregators — you pay commission on customers you already earned."),
        ("Blind service", "No single view of covers, tickets, waitlist, or guest history."),
        ("Manual grind", "Staff re-type bookings and answer the same calls every night."),
        ("No clear ROI", "Hard to see what drives spend, repeats, or empty midweek tables."),
    ]

    for i, (title, body) in enumerate(pains):
        row, col = divmod(i, 3)
        x = MARGIN + col * Inches(4.05)
        y = Inches(3.15) + row * Inches(1.85)
        add_rect(s, x, y, Inches(3.85), Inches(1.65), C["cream_soft"], line=C["line"], radius=0.06)
        add_rect(s, x, y, Inches(0.07), Inches(1.65), C["gold"])
        add_text(s, x + Inches(0.28), y + Inches(0.28), Inches(3.3), Inches(0.35), title, size=16, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
        add_text(s, x + Inches(0.28), y + Inches(0.75), Inches(3.3), Inches(0.7), body, size=13, color=C["muted"])


def slide_03_solution(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.4))
    slide_number(s, 3)

    add_text(s, MARGIN, Inches(1.15), Inches(8), Inches(0.3), "THE SOLUTION", size=11, bold=True, color=C["gold_deep"])
    gold_rule(s, MARGIN, Inches(1.5))
    add_text(
        s,
        MARGIN,
        Inches(1.7),
        Inches(12),
        Inches(0.8),
        "One platform for your whole venue.",
        size=34,
        bold=True,
        color=C["charcoal"],
        font=FONT_DISPLAY,
    )
    add_text(
        s,
        MARGIN,
        Inches(2.5),
        Inches(11),
        Inches(0.45),
        "Website, menu, ordering, reservations, customers, and analytics — in one place.",
        size=16,
        color=C["muted"],
    )

    pillars = [
        ("Premium guest website", "A branded site that matches your room — not a generic template."),
        ("Live digital & QR menu", "Publish once. Website, QR, and ordering stay in sync."),
        ("Direct online ordering", "Pickup, delivery, and dine-in on your brand — keep the margin."),
        ("Reservations that fit", "Online bookings with capacity rules, confirmations, and calendar control."),
        ("Owner dashboard", "Menus, orders, bookings, guests, and analytics together."),
        ("NZ-first by design", "Timezone, GST-aware pricing mindset, and Aotearoa hospitality pace."),
    ]

    for i, (title, body) in enumerate(pillars):
        row, col = divmod(i, 3)
        x = MARGIN + col * Inches(4.05)
        y = Inches(3.2) + row * Inches(1.85)
        add_rect(s, x, y, Inches(3.85), Inches(1.65), C["charcoal"], radius=0.06)
        add_text(
            s,
            x + Inches(0.3),
            y + Inches(0.28),
            Inches(3.25),
            Inches(0.45),
            title,
            size=15,
            bold=True,
            color=C["gold"],
            font=FONT_DISPLAY,
        )
        add_text(s, x + Inches(0.3), y + Inches(0.85), Inches(3.25), Inches(0.6), body, size=12, color=C["cream_deep"])


def slide_04_guest_journey(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.35))
    slide_number(s, 4)

    add_text(s, MARGIN, Inches(0.95), Inches(8), Inches(0.28), "GUEST JOURNEY", size=11, bold=True, color=C["gold_deep"])
    gold_rule(s, MARGIN, Inches(1.28))
    add_text(
        s,
        MARGIN,
        Inches(1.45),
        Inches(12),
        Inches(0.55),
        "From first visit to regular.",
        size=32,
        bold=True,
        color=C["charcoal"],
        font=FONT_DISPLAY,
    )

    steps = [
        ("01", "Discover", "A premium branded site guests trust."),
        ("02", "Browse", "Live menu with dietary filters and allergen notes."),
        ("03", "Order / Reserve", "Checkout or book a table — without leaving your brand."),
        ("04", "Return", "Guest history and timely follow-ups bring them back."),
    ]

    for i, (num, title, body) in enumerate(steps):
        x = MARGIN + i * Inches(3.1)
        add_rect(s, x, Inches(2.3), Inches(2.9), Inches(2.35), C["cream_soft"], line=C["line"], radius=0.06)
        add_text(s, x + Inches(0.25), Inches(2.5), Inches(2.4), Inches(0.3), num, size=12, bold=True, color=C["gold_deep"])
        add_text(
            s,
            x + Inches(0.25),
            Inches(2.95),
            Inches(2.4),
            Inches(0.4),
            title,
            size=18,
            bold=True,
            color=C["charcoal"],
            font=FONT_DISPLAY,
        )
        add_text(s, x + Inches(0.25), Inches(3.5), Inches(2.4), Inches(0.9), body, size=13, color=C["muted"])
        if i < 3:
            add_text(
                s,
                x + Inches(2.7),
                Inches(3.15),
                Inches(0.4),
                Inches(0.35),
                "→",
                size=18,
                color=C["gold"],
                align=PP_ALIGN.CENTER,
            )

    # Product strip
    imgs = [
        ASSETS / "02-harbour-kitchen.png",
        ASSETS / "02b-harbour-menu.png",
        ASSETS / "02c-harbour-order.png",
        ASSETS / "02d-harbour-reservations.png",
    ]
    for i, img in enumerate(imgs):
        if not img.exists():
            continue
        x = MARGIN + i * Inches(3.1)
        picture_frame(s, img, x, Inches(5.0), Inches(2.9), Inches(1.85))


def slide_05_direct_revenue(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.4))
    slide_number(s, 5)

    add_text(s, MARGIN, Inches(1.1), Inches(8), Inches(0.28), "DIRECT REVENUE", size=11, bold=True, color=C["gold_deep"])
    gold_rule(s, MARGIN, Inches(1.45))

    add_paras(
        s,
        MARGIN,
        Inches(1.7),
        Inches(6.2),
        Inches(1.6),
        [
            (
                "Keep customers ordering\nand booking on your brand.",
                {"size": 30, "bold": True, "color": C["charcoal"], "font": FONT_DISPLAY, "space_after": 12},
            ),
            (
                "Direct orders and bookings reduce aggregator fees on guests you already attract.",
                {"size": 15, "color": C["muted"], "space_after": 0},
            ),
        ],
    )

    points = [
        ("Convert more visitors", "Order Now and Reserve sit on a premium site guests trust."),
        ("Keep margin in-house", "Guests stay with you instead of leaking to third-party marketplaces."),
        ("Fill quiet sessions", "Live menu, offers, and analytics help shape lunch and midweek demand."),
        ("Turn covers faster", "QR menus and clear ordering cut friction from table to kitchen."),
    ]
    for i, (t, b) in enumerate(points):
        y = Inches(3.55) + i * Inches(0.8)
        add_rect(s, MARGIN, y, Inches(0.08), Inches(0.55), C["gold"])
        add_text(s, MARGIN + Inches(0.25), y, Inches(5.8), Inches(0.3), t, size=15, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
        add_text(s, MARGIN + Inches(0.25), y + Inches(0.3), Inches(5.8), Inches(0.35), b, size=13, color=C["muted"])

    img = ASSETS / "02c-harbour-order.png"
    if img.exists():
        picture_frame(s, img, Inches(7.15), Inches(1.7), Inches(5.4), Inches(4.85))


def slide_06_operations(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.4))
    slide_number(s, 6)

    add_text(s, MARGIN, Inches(1.1), Inches(8), Inches(0.28), "OPERATIONS", size=11, bold=True, color=C["gold_deep"])
    gold_rule(s, MARGIN, Inches(1.45))
    add_text(
        s,
        MARGIN,
        Inches(1.7),
        Inches(6),
        Inches(1.1),
        "One dashboard\nto run the venue.",
        size=32,
        bold=True,
        color=C["charcoal"],
        font=FONT_DISPLAY,
    )
    add_text(
        s,
        MARGIN,
        Inches(3.0),
        Inches(5.8),
        Inches(0.7),
        "Menus, orders, reservations, customers, analytics, website, QR, and team — without hopping between apps.",
        size=15,
        color=C["muted"],
    )

    ops = [
        "Live menu control — publish once, stay in sync",
        "Orders and reservations in one workspace",
        "Capacity rules you control from the dashboard",
        "Fewer tools. Fewer errors. Calmer service.",
    ]
    for i, t in enumerate(ops):
        y = Inches(3.9) + i * Inches(0.55)
        add_text(s, MARGIN, y, Inches(5.8), Inches(0.4), f"—  {t}", size=14, color=C["ink"])

    img = ASSETS / "03-dashboard.png"
    if img.exists():
        picture_frame(s, img, Inches(7.0), Inches(1.65), Inches(5.55), Inches(4.9))


def slide_07_customers(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["charcoal"])
    brand_mark(s, MARGIN, Inches(0.4), dark=True)
    slide_number(s, 7, dark=True)

    add_text(s, MARGIN, Inches(1.2), Inches(8), Inches(0.28), "CUSTOMER RELATIONSHIPS", size=11, bold=True, color=C["gold"])
    gold_rule(s, MARGIN, Inches(1.55))
    add_text(
        s,
        MARGIN,
        Inches(1.8),
        Inches(11),
        Inches(1.0),
        "Turn one-time guests into regulars.",
        size=34,
        bold=True,
        color=C["cream"],
        font=FONT_DISPLAY,
    )
    add_text(
        s,
        MARGIN,
        Inches(2.85),
        Inches(10),
        Inches(0.5),
        "Guest profiles, history, and timely follow-ups turn one-offs into people who come back.",
        size=16,
        color=C["muted_light"],
    )

    cards = [
        ("Know who walked in", "Unified guest profiles from orders and reservations — not inbox scraps."),
        ("Serve them better", "History at hand during a busy service, so regulars feel recognised."),
        ("Earn the next visit", "Own the relationship on your brand, not on someone else’s marketplace."),
    ]
    for i, (t, b) in enumerate(cards):
        x = MARGIN + i * Inches(4.05)
        add_rect(s, x, Inches(3.8), Inches(3.85), Inches(2.5), C["charcoal_soft"], radius=0.06)
        add_rect(s, x, Inches(3.8), Inches(3.85), Pt(3), C["gold"])
        add_text(
            s,
            x + Inches(0.3),
            Inches(4.15),
            Inches(3.25),
            Inches(0.55),
            t,
            size=18,
            bold=True,
            color=C["gold_soft"],
            font=FONT_DISPLAY,
        )
        add_text(s, x + Inches(0.3), Inches(4.85), Inches(3.25), Inches(1.1), b, size=14, color=C["cream_deep"])


def slide_08_analytics(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.4))
    slide_number(s, 8)

    add_text(s, MARGIN, Inches(1.1), Inches(8), Inches(0.28), "ANALYTICS", size=11, bold=True, color=C["gold_deep"])
    gold_rule(s, MARGIN, Inches(1.45))
    add_text(
        s,
        MARGIN,
        Inches(1.7),
        Inches(6),
        Inches(1.0),
        "Stop guessing.\nStart seeing.",
        size=34,
        bold=True,
        color=C["charcoal"],
        font=FONT_DISPLAY,
    )
    add_text(
        s,
        MARGIN,
        Inches(3.0),
        Inches(5.8),
        Inches(0.7),
        "Revenue, AOV, bestsellers, reservations, and conversion funnel — facts from your venue data. Never invented numbers.",
        size=15,
        color=C["muted"],
    )

    metrics = [
        ("Revenue", "What the room actually made"),
        ("AOV", "How tickets are building"),
        ("Bestsellers", "What to push — and protect"),
        ("Funnel", "Where guests drop off"),
    ]
    for i, (t, b) in enumerate(metrics):
        y = Inches(3.9) + i * Inches(0.7)
        add_text(s, MARGIN, y, Inches(1.6), Inches(0.35), t, size=14, bold=True, color=C["gold_deep"], font=FONT_DISPLAY)
        add_text(s, MARGIN + Inches(1.7), y, Inches(4), Inches(0.35), b, size=14, color=C["ink"])

    img = ASSETS / "03-dashboard-analytics.png"
    if img.exists():
        picture_frame(s, img, Inches(7.0), Inches(1.65), Inches(5.55), Inches(4.9))


def slide_09_before_after(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.35))
    slide_number(s, 9)

    add_text(s, MARGIN, Inches(1.0), Inches(8), Inches(0.28), "THE SHIFT", size=11, bold=True, color=C["gold_deep"])
    gold_rule(s, MARGIN, Inches(1.35))
    add_text(
        s,
        MARGIN,
        Inches(1.55),
        Inches(11),
        Inches(0.55),
        "Before vs after RR NOVA.",
        size=32,
        bold=True,
        color=C["charcoal"],
        font=FONT_DISPLAY,
    )

    before = [
        "5–10 logins every service",
        "Stale menus & broken booking links",
        "Aggregator fees eating margin",
        "Guest data trapped in inboxes",
        "Managers flying blind on covers",
        "Marketing is an afterthought",
    ]
    after = [
        "One dashboard for the whole venue",
        "Live menu, QR, and site in sync",
        "Direct orders with secure checkout",
        "Unified guest profiles & history",
        "Live boards for kitchen & floor",
        "Analytics guiding real decisions",
    ]

    # Before panel
    add_rect(s, MARGIN, Inches(2.4), Inches(5.7), Inches(4.4), C["cream_soft"], line=C["line"], radius=0.06)
    add_text(
        s,
        MARGIN + Inches(0.4),
        Inches(2.65),
        Inches(4.8),
        Inches(0.4),
        "Before",
        size=20,
        bold=True,
        color=C["danger_soft"],
        font=FONT_DISPLAY,
    )
    for i, t in enumerate(before):
        add_text(
            s,
            MARGIN + Inches(0.4),
            Inches(3.25) + i * Inches(0.5),
            Inches(5.0),
            Inches(0.4),
            f"✕   {t}",
            size=14,
            color=C["ink"],
        )

    # After panel
    add_rect(s, Inches(6.95), Inches(2.4), Inches(5.7), Inches(4.4), C["charcoal"], radius=0.06)
    add_text(
        s,
        Inches(7.35),
        Inches(2.65),
        Inches(4.8),
        Inches(0.4),
        "After RR NOVA",
        size=20,
        bold=True,
        color=C["gold"],
        font=FONT_DISPLAY,
    )
    for i, t in enumerate(after):
        add_text(
            s,
            Inches(7.35),
            Inches(3.25) + i * Inches(0.5),
            Inches(5.0),
            Inches(0.4),
            f"✓   {t}",
            size=14,
            color=C["cream"],
        )


def slide_10_cta(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    solid_bg(s, C["charcoal"])
    add_rect(s, 0, 0, Inches(0.08), SLIDE_H, C["gold"])
    brand_mark(s, MARGIN, Inches(0.45), dark=True)
    slide_number(s, 10, dark=True)

    add_text(
        s,
        MARGIN,
        Inches(2.0),
        Inches(11),
        Inches(0.35),
        "NEXT STEP",
        size=12,
        bold=True,
        color=C["gold"],
    )
    add_text(
        s,
        MARGIN,
        Inches(2.5),
        Inches(11.5),
        Inches(1.3),
        "What could this look like\nfor your restaurant?",
        size=38,
        bold=True,
        color=C["cream"],
        font=FONT_DISPLAY,
    )
    add_text(
        s,
        MARGIN,
        Inches(4.15),
        Inches(10),
        Inches(0.7),
        "Explore the Harbour Kitchen live demo — or start your own workspace.",
        size=16,
        color=C["muted_light"],
    )

    # CTA pill
    add_rect(s, MARGIN, Inches(5.1), Inches(5.2), Inches(0.65), C["gold"], radius=0.12)
    add_text(
        s,
        MARGIN,
        Inches(5.22),
        Inches(5.2),
        Inches(0.45),
        DEMO_URL,
        size=14,
        bold=True,
        color=C["charcoal"],
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    add_text(
        s,
        MARGIN,
        Inches(6.15),
        Inches(11),
        Inches(0.35),
        "Platform home  ·  Guest site /r/harbour-kitchen  ·  Owner dashboard /dashboard/menus",
        size=12,
        color=C["muted"],
    )
    add_text(
        s,
        MARGIN,
        Inches(6.7),
        Inches(11),
        Inches(0.3),
        "RR NOVA  ·  Your restaurant. Your customers. Your data. One platform.",
        size=12,
        color=C["gold_soft"],
        italic=True,
    )


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_01_cover(prs)
    slide_02_problem(prs)
    slide_03_solution(prs)
    slide_04_guest_journey(prs)
    slide_05_direct_revenue(prs)
    slide_06_operations(prs)
    slide_07_customers(prs)
    slide_08_analytics(prs)
    slide_09_before_after(prs)
    slide_10_cta(prs)

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"Wrote {OUT_PPTX} ({len(prs.slides)} slides)")
    return OUT_PPTX


if __name__ == "__main__":
    build()
