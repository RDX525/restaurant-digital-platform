#!/usr/bin/env python3
"""Generate a premium 10-slide RR NOVA restaurant-owner sales deck (16:9 PPTX)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

# ── Design tokens — restrained hospitality palette ──────────────────────────
C = {
    "cream": RGBColor(0xF7, 0xF3, 0xEC),
    "cream2": RGBColor(0xFB, 0xF8, 0xF3),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "charcoal": RGBColor(0x1C, 0x1A, 0x17),
    "ink": RGBColor(0x2A, 0x27, 0x23),
    "muted": RGBColor(0x7A, 0x73, 0x68),
    "muted_soft": RGBColor(0xA8, 0xA0, 0x94),
    "line": RGBColor(0xE6, 0xDF, 0xD4),
    "gold": RGBColor(0xB8, 0x9B, 0x6A),
    "gold_soft": RGBColor(0xD4, 0xC4, 0xA0),
    "dark": RGBColor(0x14, 0x13, 0x11),
    "dark2": RGBColor(0x22, 0x20, 0x1C),
    "panel": RGBColor(0x2C, 0x29, 0x24),
    "before_bg": RGBColor(0xF0, 0xEA, 0xE2),
    "after_bg": RGBColor(0x14, 0x13, 0x11),
    "danger": RGBColor(0xA8, 0x5A, 0x4A),
}

FONT_DISPLAY = "Georgia"
FONT_SANS = "Avenir Next"
FONT_FALLBACK = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.65)


def _set_run_font(run, size, bold=False, color=None, font=FONT_SANS, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color


def _fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _line(shape, color: RGBColor | None = None, width_pt: float = 0.75):
    if color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color
        shape.line.width = Pt(width_pt)


def _shadow(shape, soft=True):
    spPr = shape._element.spPr
    effectLst = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    if effectLst is None:
        effectLst = OxmlElement("a:effectLst")
        spPr.append(effectLst)
    for child in list(effectLst):
        effectLst.remove(child)
    outer = OxmlElement("a:outerShdw")
    outer.set("blurRad", "50800" if soft else "76200")
    outer.set("dist", "25400")
    outer.set("dir", "2700000")
    outer.set("algn", "tl")
    outer.set("rotWithShape", "0")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", "1C1A17")
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "12000")
    srgb.append(alpha)
    outer.append(srgb)
    effectLst.append(outer)


def rect(slide, left, top, width, height, fill, line=None, radius=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    _fill(shp, fill)
    _line(shp, line)
    if radius is not None and hasattr(shp, "adjustments") and len(shp.adjustments) > 0:
        shp.adjustments[0] = radius
    if shadow:
        _shadow(shp)
    return shp


def txt(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=14,
    bold=False,
    color=C["ink"],
    font=FONT_SANS,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    italic=False,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, bold=bold, color=color, font=font, italic=italic)
    return box


def paras(slide, left, top, width, height, lines: list[tuple[str, dict]]):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, style) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = style.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(style.get("space_before", 0))
        p.space_after = Pt(style.get("space_after", 2))
        run = p.add_run()
        run.text = text
        _set_run_font(
            run,
            style.get("size", 14),
            bold=style.get("bold", False),
            color=style.get("color", C["ink"]),
            font=style.get("font", FONT_SANS),
            italic=style.get("italic", False),
        )
    return box


def bg(slide, color):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, color)


def gold_rule(slide, left, top, width=Inches(1.2)):
    rect(slide, left, top, width, Inches(0.03), C["gold"])


def brand_mark(slide, left, top, dark=False):
    """Small RR NOVA mark — gold square + wordmark."""
    rect(slide, left, top, Inches(0.28), Inches(0.28), C["gold"], radius=0.08)
    txt(
        slide,
        left,
        top + Inches(0.02),
        Inches(0.28),
        Inches(0.26),
        "RR",
        size=8,
        bold=True,
        color=C["dark"],
        align=PP_ALIGN.CENTER,
    )
    txt(
        slide,
        left + Inches(0.36),
        top + Inches(0.02),
        Inches(1.5),
        Inches(0.28),
        "RR NOVA",
        size=11,
        bold=True,
        color=C["cream"] if dark else C["charcoal"],
        font=FONT_SANS,
    )


def slide_num(slide, n, total=10, dark=False):
    txt(
        slide,
        SLIDE_W - Inches(1.4),
        SLIDE_H - Inches(0.45),
        Inches(0.9),
        Inches(0.3),
        f"{n:02d} / {total:02d}",
        size=10,
        color=C["muted_soft"] if not dark else RGBColor(0x6A, 0x64, 0x5A),
        align=PP_ALIGN.RIGHT,
    )


def footer_brand(slide, dark=False):
    brand_mark(slide, MARGIN, SLIDE_H - Inches(0.5), dark=dark)


# ── Slides ──────────────────────────────────────────────────────────────────


def slide_01_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["dark"])
    # soft editorial panels
    rect(s, Inches(8.8), Inches(-0.8), Inches(5.5), Inches(4.2), C["dark2"], radius=0.2)
    rect(s, Inches(-1.2), Inches(5.2), Inches(5), Inches(3.2), C["panel"], radius=0.2)

    txt(s, MARGIN, Inches(0.55), Inches(3), Inches(0.3), "RR NOVA", size=12, bold=True, color=C["gold"])
    gold_rule(s, MARGIN, Inches(0.95), Inches(0.9))

    paras(
        s,
        MARGIN,
        Inches(1.5),
        Inches(8),
        Inches(2.4),
        [
            ("Your restaurant.", {"size": 40, "bold": True, "color": C["cream"], "font": FONT_DISPLAY, "space_after": 0}),
            ("Your customers.", {"size": 40, "bold": True, "color": C["cream"], "font": FONT_DISPLAY, "space_after": 0}),
            ("Your data.", {"size": 40, "bold": True, "color": C["gold_soft"], "font": FONT_DISPLAY, "space_after": 0}),
        ],
    )

    txt(
        s,
        MARGIN,
        Inches(4.35),
        Inches(7.2),
        Inches(1.1),
        "One platform to attract customers, take direct orders,\nmanage reservations, and understand what drives your business.",
        size=15,
        color=C["muted_soft"],
    )

    # Editorial journey — not a flowchart
    steps = ["DISCOVER", "MENU", "ORDER / RESERVE", "CUSTOMER", "REPEAT"]
    x0 = MARGIN
    y = Inches(6.05)
    for i, step in enumerate(steps):
        txt(s, x0, y, Inches(2.0), Inches(0.28), step, size=10, bold=True, color=C["gold"])
        if i < len(steps) - 1:
            txt(s, x0 + Inches(1.85), y, Inches(0.35), Inches(0.28), "→", size=12, color=C["muted_soft"])
            x0 += Inches(2.2)

    slide_num(s, 1, dark=True)


def slide_02_problem(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(11), Inches(0.9), "Your restaurant shouldn't need\n10 different systems.", size=30, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN,
        Inches(2.15),
        Inches(11),
        Inches(0.45),
        "Website, menu, bookings, ordering, spreadsheets and customer information often live in separate places.",
        size=14,
        color=C["muted"],
    )

    problems = [
        ("Too many tools", "Staff jump between systems and information doesn't talk to each other."),
        ("Menu chaos", "Printed and third-party menus can become stale when dishes change."),
        ("Missed revenue", "Customers can be pushed to aggregators where commissions reduce margin."),
        ("Manual work", "Teams spend time re-entering bookings, chasing confirmations and answering calls."),
        ("Limited visibility", "Owners struggle to see covers, tickets, guest history and performance together."),
        ("Unclear ROI", "It can be difficult to see what drives spend, repeat visits and quiet sessions."),
    ]
    for i, (title, body) in enumerate(problems):
        row, col = divmod(i, 3)
        x = MARGIN + col * Inches(4.05)
        y = Inches(2.85) + row * Inches(2.05)
        rect(s, x, y, Inches(3.85), Inches(1.85), C["white"], line=C["line"], radius=0.04, shadow=True)
        rect(s, x, y, Inches(0.07), Inches(1.85), C["gold"])
        txt(s, x + Inches(0.28), y + Inches(0.28), Inches(3.3), Inches(0.35), title, size=15, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
        txt(s, x + Inches(0.28), y + Inches(0.75), Inches(3.3), Inches(0.85), body, size=12, color=C["muted"])

    slide_num(s, 2)
    footer_brand(s)


def slide_03_solution(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream2"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(10), Inches(0.55), "One operating system for your venue.", size=30, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN,
        Inches(1.7),
        Inches(10),
        Inches(0.4),
        "RR NOVA connects the guest experience and the owner's operations in one place.",
        size=14,
        color=C["muted"],
    )

    # Center hub
    cx, cy = Inches(5.55), Inches(4.05)
    rect(s, cx, cy, Inches(2.2), Inches(1.0), C["dark"], radius=0.06)
    txt(s, cx, cy + Inches(0.32), Inches(2.2), Inches(0.4), "RR NOVA", size=16, bold=True, color=C["gold"], align=PP_ALIGN.CENTER)

    caps = [
        (Inches(0.7), Inches(2.55), "Website", "Premium branded\nguest experience."),
        (Inches(3.5), Inches(2.55), "Live menu", "Website + QR + ordering\nstay in sync."),
        (Inches(8.2), Inches(2.55), "Direct ordering", "Pickup, delivery and\ndine-in on your brand."),
        (Inches(10.7), Inches(2.55), "Reservations", "Capacity rules,\nconfirmations, calendar."),
        (Inches(2.0), Inches(5.55), "Owner dashboard", "Menus, orders, bookings,\nguests and team."),
        (Inches(9.2), Inches(5.55), "Analytics", "Revenue, AOV, bestsellers\nand conversion funnel."),
    ]
    for x, y, title, body in caps:
        rect(s, x, y, Inches(2.35), Inches(1.25), C["white"], line=C["line"], radius=0.04, shadow=True)
        txt(s, x + Inches(0.18), y + Inches(0.2), Inches(2.0), Inches(0.3), title, size=13, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
        txt(s, x + Inches(0.18), y + Inches(0.55), Inches(2.0), Inches(0.6), body, size=11, color=C["muted"])

    slide_num(s, 3)
    footer_brand(s)


def slide_04_journey(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(11), Inches(0.5), "Keep the customer in your brand.", size=30, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN,
        Inches(1.7),
        Inches(11),
        Inches(0.35),
        "From discovery to booking or ordering, RR NOVA creates a direct path.",
        size=14,
        color=C["muted"],
    )

    stages = [
        ("01", "DISCOVER", "Branded restaurant website."),
        ("02", "BROWSE", "Live menu with dietary and allergen information."),
        ("03", "ACT", "Order directly or reserve a table."),
        ("04", "RETURN", "Use guest history and follow-ups to encourage repeat visits."),
    ]
    for i, (num, title, body) in enumerate(stages):
        x = MARGIN + i * Inches(3.1)
        rect(s, x, Inches(2.5), Inches(2.9), Inches(3.6), C["white"], line=C["line"], radius=0.04, shadow=True)
        txt(s, x + Inches(0.25), Inches(2.8), Inches(2.4), Inches(0.3), num, size=12, bold=True, color=C["gold"])
        txt(s, x + Inches(0.25), Inches(3.3), Inches(2.4), Inches(0.45), title, size=16, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
        gold_rule(s, x + Inches(0.25), Inches(3.9), Inches(0.55))
        txt(s, x + Inches(0.25), Inches(4.2), Inches(2.4), Inches(1.3), body, size=13, color=C["muted"])
        if i < 3:
            txt(s, x + Inches(2.75), Inches(4.0), Inches(0.4), Inches(0.35), "→", size=18, color=C["gold"])

    slide_num(s, 4)
    footer_brand(s)


def slide_05_direct_revenue(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream2"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(12), Inches(0.55), "Why send a customer somewhere else to order?", size=28, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN,
        Inches(1.7),
        Inches(12),
        Inches(0.4),
        "RR NOVA puts direct ordering and reservations on the restaurant's own branded experience.",
        size=14,
        color=C["muted"],
    )

    # Left — ordering
    rect(s, MARGIN, Inches(2.35), Inches(5.85), Inches(3.55), C["white"], line=C["line"], radius=0.04, shadow=True)
    txt(s, MARGIN + Inches(0.4), Inches(2.6), Inches(5), Inches(0.35), "Direct ordering", size=18, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN + Inches(0.4),
        Inches(3.1),
        Inches(5),
        Inches(0.35),
        "Keep the ordering experience on your brand.",
        size=13,
        italic=True,
        color=C["gold"],
        font=FONT_DISPLAY,
    )
    for i, item in enumerate(["Pickup", "Delivery", "Dine-in"]):
        y = Inches(3.7) + i * Inches(0.55)
        rect(s, MARGIN + Inches(0.4), y, Inches(0.18), Inches(0.18), C["gold"], radius=0.5)
        txt(s, MARGIN + Inches(0.75), y - Inches(0.02), Inches(4), Inches(0.3), item, size=14, color=C["ink"])
    txt(
        s,
        MARGIN + Inches(0.4),
        Inches(5.35),
        Inches(5),
        Inches(0.4),
        "Reduce dependence on third-party marketplaces\nand retain more of each order.",
        size=12,
        color=C["muted"],
    )

    # Right — reservations
    rect(s, Inches(6.9), Inches(2.35), Inches(5.85), Inches(3.55), C["dark"], radius=0.04, shadow=True)
    txt(s, Inches(7.3), Inches(2.6), Inches(5), Inches(0.35), "Direct reservations", size=18, bold=True, color=C["cream"], font=FONT_DISPLAY)
    txt(s, Inches(7.3), Inches(3.15), Inches(5), Inches(0.3), "Guests choose", size=11, bold=True, color=C["gold"])
    for i, item in enumerate(["Party size", "Date", "Time"]):
        txt(s, Inches(7.3), Inches(3.5) + i * Inches(0.35), Inches(5), Inches(0.3), f"·  {item}", size=13, color=C["cream"])
    txt(s, Inches(7.3), Inches(4.65), Inches(5), Inches(0.3), "You control", size=11, bold=True, color=C["gold"])
    for i, item in enumerate(["Capacity rules", "Confirmations", "Calendar flow"]):
        txt(s, Inches(7.3), Inches(5.0) + i * Inches(0.32), Inches(5), Inches(0.3), f"·  {item}", size=13, color=C["cream"])

    txt(
        s,
        MARGIN,
        Inches(6.2),
        Inches(12),
        Inches(0.35),
        "SELL DIRECTLY   ·   KEEP CONTROL   ·   OWN THE EXPERIENCE",
        size=12,
        bold=True,
        color=C["charcoal"],
        align=PP_ALIGN.CENTER,
    )

    slide_num(s, 5)
    footer_brand(s)


def slide_06_operations(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(8), Inches(0.55), "Give your team one place to run service.", size=28, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN,
        Inches(1.7),
        Inches(8),
        Inches(0.4),
        "Less hopping between apps. More visibility during a busy shift.",
        size=14,
        color=C["muted"],
    )

    # Dashboard mock — editorial, not fake metrics
    frame_x, frame_y = Inches(0.65), Inches(2.35)
    frame_w, frame_h = Inches(12.0), Inches(4.35)
    rect(s, frame_x, frame_y, frame_w, frame_h, C["dark"], radius=0.05, shadow=True)

    # Sidebar
    rect(s, frame_x, frame_y, Inches(2.6), frame_h, C["dark2"])
    txt(s, frame_x + Inches(0.25), frame_y + Inches(0.25), Inches(2), Inches(0.25), "WORKSPACE", size=9, bold=True, color=C["gold"])
    nav = ["Menus", "Orders", "Reservations", "Customers", "Website", "QR", "Analytics", "Team"]
    for i, item in enumerate(nav):
        y = frame_y + Inches(0.65) + i * Inches(0.38)
        if i == 0:
            rect(s, frame_x + Inches(0.15), y - Inches(0.05), Inches(2.3), Inches(0.34), C["panel"], radius=0.04)
        txt(s, frame_x + Inches(0.3), y, Inches(2), Inches(0.28), item, size=12, color=C["cream"] if i == 0 else C["muted_soft"])

    # Main content
    txt(s, frame_x + Inches(2.95), frame_y + Inches(0.35), Inches(5), Inches(0.25), "OWNER WORKSPACE", size=10, bold=True, color=C["gold"])
    txt(s, frame_x + Inches(2.95), frame_y + Inches(0.7), Inches(8), Inches(0.45), "One dashboard for the whole venue.", size=22, bold=True, color=C["cream"], font=FONT_DISPLAY)
    txt(
        s,
        frame_x + Inches(2.95),
        frame_y + Inches(1.35),
        Inches(8.5),
        Inches(0.5),
        "Menus, orders, reservations, customers, website, QR, analytics and team —\nin one calm workspace built for service.",
        size=13,
        color=C["muted_soft"],
    )

    tiles = [
        ("Menus", "Keep dishes live everywhere"),
        ("Orders", "Track the live queue"),
        ("Reservations", "Confirm and seat with control"),
        ("Customers", "See profiles & history"),
    ]
    for i, (t, b) in enumerate(tiles):
        x = frame_x + Inches(2.95) + i * Inches(2.2)
        rect(s, x, frame_y + Inches(2.3), Inches(2.05), Inches(1.5), C["panel"], radius=0.04)
        txt(s, x + Inches(0.15), frame_y + Inches(2.55), Inches(1.75), Inches(0.3), t, size=13, bold=True, color=C["gold"], font=FONT_DISPLAY)
        txt(s, x + Inches(0.15), frame_y + Inches(3.0), Inches(1.75), Inches(0.55), b, size=11, color=C["muted_soft"])

    slide_num(s, 6)
    footer_brand(s)


def slide_07_customers(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream2"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(12), Inches(0.9), "The first order shouldn't be\nthe end of the relationship.", size=28, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN,
        Inches(2.15),
        Inches(12),
        Inches(0.4),
        "Guest profiles and history can help restaurants understand customers and create timely follow-ups.",
        size=14,
        color=C["muted"],
    )

    sections = [
        ("KNOW THE GUEST", "Unified profiles and visit/order history give the restaurant a clearer picture of its customers."),
        ("SPOT PATTERNS", "Understand what customers buy, how often they return, and where demand is coming from."),
        ("BRING THEM BACK", "Use timely follow-ups and customer knowledge to turn one-offs into regulars."),
    ]
    for i, (title, body) in enumerate(sections):
        x = MARGIN + i * Inches(4.1)
        rect(s, x, Inches(2.9), Inches(3.9), Inches(3.2), C["white"], line=C["line"], radius=0.04, shadow=True)
        txt(s, x + Inches(0.35), Inches(3.25), Inches(3.2), Inches(0.25), f"0{i+1}", size=11, bold=True, color=C["gold"])
        txt(s, x + Inches(0.35), Inches(3.7), Inches(3.2), Inches(0.7), title, size=16, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
        gold_rule(s, x + Inches(0.35), Inches(4.55), Inches(0.6))
        txt(s, x + Inches(0.35), Inches(4.8), Inches(3.2), Inches(1.0), body, size=13, color=C["muted"])

    slide_num(s, 7)
    footer_brand(s)


def slide_08_analytics(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(10), Inches(0.5), "Stop guessing. Start seeing.", size=30, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
    txt(
        s,
        MARGIN,
        Inches(1.7),
        Inches(11),
        Inches(0.4),
        "RR NOVA brings key venue metrics into one view.",
        size=14,
        color=C["muted"],
    )

    metrics = [
        ("Revenue", "How much the venue is generating."),
        ("AOV", "Average order value."),
        ("Bestsellers", "What customers actually buy."),
        ("Reservations", "Booking activity and demand."),
        ("Conversion", "How visitors turn into actions."),
    ]
    for i, (title, body) in enumerate(metrics):
        x = MARGIN + i * Inches(2.45)
        rect(s, x, Inches(2.4), Inches(2.3), Inches(2.55), C["white"], line=C["line"], radius=0.04, shadow=True)
        # UI illustration bars — clearly decorative, no numbers claimed as results
        for j, h in enumerate([0.55, 0.85, 0.45, 0.7, 0.95]):
            bar_h = Inches(h * 0.55)
            rect(
                s,
                x + Inches(0.25) + j * Inches(0.35),
                Inches(3.55) - bar_h + Inches(0.55),
                Inches(0.22),
                bar_h,
                C["gold_soft"] if j != 4 else C["gold"],
                radius=0.02,
            )
        txt(s, x + Inches(0.2), Inches(3.85), Inches(1.9), Inches(0.3), title, size=13, bold=True, color=C["charcoal"], font=FONT_DISPLAY)
        txt(s, x + Inches(0.2), Inches(4.25), Inches(1.9), Inches(0.55), body, size=11, color=C["muted"])

    rect(s, MARGIN, Inches(5.3), Inches(12.0), Inches(1.15), C["dark"], radius=0.04)
    txt(
        s,
        MARGIN + Inches(0.5),
        Inches(5.65),
        Inches(11),
        Inches(0.5),
        "Use venue data to shape demand, understand quiet sessions and make real decisions.",
        size=15,
        color=C["cream"],
        align=PP_ALIGN.CENTER,
        font=FONT_DISPLAY,
    )

    slide_num(s, 8)
    footer_brand(s)


def slide_09_before_after(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["cream2"])
    brand_mark(s, MARGIN, Inches(0.35))
    gold_rule(s, MARGIN, Inches(0.85), Inches(0.7))
    txt(s, MARGIN, Inches(1.05), Inches(12), Inches(0.5), "From fragmented tools to one connected venue.", size=26, bold=True, color=C["charcoal"], font=FONT_DISPLAY)

    # Before — fragmented feel
    rect(s, MARGIN, Inches(1.85), Inches(5.85), Inches(4.85), C["before_bg"], radius=0.04)
    txt(s, MARGIN + Inches(0.4), Inches(2.15), Inches(5), Inches(0.4), "BEFORE", size=14, bold=True, color=C["danger"])
    before = [
        "5–10 logins every service",
        "Stale menus & broken booking links",
        "Aggregator fees eating margin",
        "Guest data trapped in inboxes",
        "Managers flying blind on covers",
        "Marketing is an afterthought",
    ]
    for i, line in enumerate(before):
        # small scattered chips to feel fragmented
        y = Inches(2.75) + i * Inches(0.55)
        rect(s, MARGIN + Inches(0.4), y, Inches(5.05), Inches(0.45), C["white"], line=C["line"], radius=0.03)
        txt(s, MARGIN + Inches(0.55), y + Inches(0.1), Inches(4.7), Inches(0.3), f"✕   {line}", size=13, color=C["ink"])

    # After — calm & controlled
    rect(s, Inches(6.9), Inches(1.85), Inches(5.85), Inches(4.85), C["after_bg"], radius=0.04)
    txt(s, Inches(7.3), Inches(2.15), Inches(5), Inches(0.4), "AFTER RR NOVA", size=14, bold=True, color=C["gold"])
    after = [
        "One dashboard for the whole venue",
        "Live menu, QR and site in sync",
        "Direct orders with secure checkout",
        "Unified guest profiles & history",
        "Live boards for kitchen & floor",
        "Analytics guiding real decisions",
    ]
    for i, line in enumerate(after):
        txt(s, Inches(7.3), Inches(2.85) + i * Inches(0.5), Inches(5.1), Inches(0.4), f"✓   {line}", size=14, color=C["cream"])

    slide_num(s, 9)
    footer_brand(s)


def slide_10_cta(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C["dark"])
    rect(s, Inches(8.5), Inches(-1), Inches(6), Inches(4), C["dark2"], radius=0.25)
    rect(s, Inches(-1), Inches(5.5), Inches(4.5), Inches(3), C["panel"], radius=0.25)

    txt(s, MARGIN, Inches(0.55), Inches(3), Inches(0.3), "RR NOVA", size=12, bold=True, color=C["gold"])
    gold_rule(s, MARGIN, Inches(0.95), Inches(0.9))

    txt(
        s,
        MARGIN,
        Inches(1.8),
        Inches(11),
        Inches(1.1),
        "What could this look like\nfor your restaurant?",
        size=34,
        bold=True,
        color=C["cream"],
        font=FONT_DISPLAY,
    )
    txt(
        s,
        MARGIN,
        Inches(3.5),
        Inches(10),
        Inches(0.8),
        "Let's look at your current website, ordering and booking journey—\nand show you where RR NOVA can simplify the experience.",
        size=15,
        color=C["muted_soft"],
    )

    rect(s, MARGIN, Inches(4.7), Inches(2.6), Inches(0.55), C["gold"], radius=0.04)
    txt(s, MARGIN, Inches(4.82), Inches(2.6), Inches(0.35), "EXPLORE RR NOVA", size=13, bold=True, color=C["dark"], align=PP_ALIGN.CENTER)

    txt(s, MARGIN + Inches(2.9), Inches(4.85), Inches(6), Inches(0.35), "See the Harbour Kitchen live demo", size=14, color=C["cream"])
    txt(
        s,
        MARGIN,
        Inches(5.55),
        Inches(10),
        Inches(0.3),
        "restaurant-digital-platform-beige.vercel.app",
        size=12,
        color=C["gold_soft"],
    )

    txt(
        s,
        MARGIN,
        SLIDE_H - Inches(0.55),
        Inches(8),
        Inches(0.3),
        "RR NOVA  ·  Crafted for Aotearoa New Zealand",
        size=11,
        color=C["muted_soft"],
    )
    slide_num(s, 10, dark=True)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_01_cover(prs)
    slide_02_problem(prs)
    slide_03_solution(prs)
    slide_04_journey(prs)
    slide_05_direct_revenue(prs)
    slide_06_operations(prs)
    slide_07_customers(prs)
    slide_08_analytics(prs)
    slide_09_before_after(prs)
    slide_10_cta(prs)

    out_docs = Path(__file__).resolve().parents[1] / "docs" / "presentations" / "RR-NOVA-Restaurant-Owner-Sales-Deck.pptx"
    out_desktop = Path.home() / "Desktop" / "RR-NOVA-Restaurant.pptx"
    out_docs.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_docs))
    prs.save(str(out_desktop))
    print(f"Wrote {out_docs}")
    print(f"Wrote {out_desktop}")
    print(f"{len(prs.slides)} slides")
    return out_desktop


if __name__ == "__main__":
    build()
